import os
import csv
from openpyxl import load_workbook


# ----------------------------------------------------------
# File reading helpers
# ----------------------------------------------------------
def get_file_headers(filename):
    """Read the header row from .xlsx/.xlsm/.csv files."""
    if filename.lower().endswith(('.xlsx', '.xlsm')):
        wb = load_workbook(filename, read_only=True)
        ws = wb.active
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [str(h) if h is not None else f"Column {i+1}"
                   for i, h in enumerate(first_row)]
        wb.close()
        return headers
    elif filename.lower().endswith('.csv'):
        with open(filename, 'r', encoding='utf-8-sig') as f:
            return next(csv.reader(f))
    raise ValueError("Unsupported file type. Use .xlsx, .xlsm, or .csv")


def _read_rows(filename):
    """Return (headers, list_of_rows_as_tuples)."""
    if filename.lower().endswith(('.xlsx', '.xlsm')):
        wb = load_workbook(filename, data_only=True)
        ws = wb.active
        data = list(ws.iter_rows(values_only=True))
        wb.close()
        if not data:
            return [], []
        headers = [str(c) if c is not None else "" for c in data[0]]
        return headers, data[1:]
    elif filename.lower().endswith('.csv'):
        with open(filename, 'r', encoding='utf-8-sig') as f:
            data = list(csv.reader(f))
        if not data:
            return [], []
        return data[0], data[1:]
    raise ValueError("Unsupported file type")


def _norm(val):
    return str(val).strip().lower() if val is not None else ""


def is_empty(val):
    """Same semantics as DataCleanerTool.is_empty() in the dashboard."""
    if val is None:
        return True
    if isinstance(val, str) and val.strip() == "":
        return True
    return False


# ----------------------------------------------------------
# Main processing
# ----------------------------------------------------------
def remove_coordinators(
    filename,
    # --- coordinator rules ---
    mode="role_column",              # "role_column" | "manual_list" | "both"
    column=None,                     # name of role/designation column
    keywords=None,                   # list of keywords for role match
    manual_list=None,                # list of names/emails/phones/ids
    # --- empty-row rule ---
    remove_empty=False,              # bool: enable empty-cell removal
    empty_column=None,               # column name OR "ALL COLUMNS (Check entire row)"
):
    """
    Removes rows that match coordinator rules AND/OR have empty cells.

    Returns:
        headers          : list[str]
        kept_rows        : list[tuple]
        removed_rows     : list[tuple]
        removed_reasons  : list[str]   # parallel to removed_rows
        stats            : dict        # {'empty': n, 'coordinator': n}
    """
    headers, rows = _read_rows(filename)
    if not headers:
        return [], [], [], [], {"empty": 0, "coordinator": 0}

    keywords = [k.strip().lower() for k in (keywords or []) if k.strip()]
    manual_list = [m.strip().lower() for m in (manual_list or []) if m.strip()]

    # --- locate columns for manual-list matching ---
    lower_headers = [h.lower() for h in headers]
    match_cols = []
    for candidate in ("name", "email", "phone", "idnumber", "id number", "id"):
        for i, h in enumerate(lower_headers):
            if candidate in h:
                match_cols.append(i)
                break
    if not match_cols:
        match_cols = list(range(len(headers)))

    # --- role column index ---
    role_col_idx = headers.index(column) if column and column in headers else None

    # --- empty column index / mode ---
    empty_col_idx = None
    empty_all_cols = False
    if remove_empty:
        if empty_column == "ALL COLUMNS (Check entire row)" or not empty_column:
            empty_all_cols = True
        elif empty_column in headers:
            empty_col_idx = headers.index(empty_column)

    kept_rows, removed_rows, removed_reasons = [], [], []
    stats = {"empty": 0, "coordinator": 0}

    for row in rows:
        reason = None

        # ---- RULE A: empty cells ----
        if remove_empty:
            if empty_all_cols:
                if any(is_empty(v) for v in row):
                    reason = "Empty value in row"
                    stats["empty"] += 1
            elif empty_col_idx is not None:
                val = row[empty_col_idx] if empty_col_idx < len(row) else None
                if is_empty(val):
                    reason = f"Empty value in column '{empty_column}'"
                    stats["empty"] += 1

        # ---- RULE B: coordinator rules ----
        if reason is None and mode in ("role_column", "both") and role_col_idx is not None:
            val = _norm(row[role_col_idx]) if role_col_idx < len(row) else ""
            if any(kw in val for kw in keywords):
                reason = f"Role column '{column}' matched a coordinator keyword"
                stats["coordinator"] += 1

        if reason is None and mode in ("manual_list", "both") and manual_list:
            cells = [_norm(row[i]) if i < len(row) else "" for i in match_cols]
            for cell in cells:
                if cell and any(entry in cell or cell == entry for entry in manual_list):
                    reason = "Matched manual coordinator list"
                    stats["coordinator"] += 1
                    break

        if reason:
            removed_rows.append(row)
            removed_reasons.append(reason)
        else:
            kept_rows.append(row)

    return headers, kept_rows, removed_rows, removed_reasons, stats