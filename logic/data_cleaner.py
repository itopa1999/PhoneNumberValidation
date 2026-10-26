import csv
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter

def get_file_headers(filepath):
    """Reads the first row of the CSV or Excel file to get the headers."""
    if filepath.lower().endswith(('.xlsx', '.xlsm')):
        wb = load_workbook(filepath, read_only=True)
        ws = wb.active
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [str(h) if h is not None else f"Column {i+1}" for i, h in enumerate(first_row)]
        wb.close()
        return headers
    elif filepath.lower().endswith('.csv'):
        with open(filepath, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            return next(reader)
    return []

def is_empty(val):
    """Helper to determine if a cell is truly empty."""
    if val is None:
        return True
    if isinstance(val, str) and val.strip() == "":
        return True
    return False

def process_file_data(filepath, col_name):
    """
    Reads the file and splits rows into kept and removed.
    Returns: headers, kept_rows, removed_rows
    """
    headers = []
    kept_rows = []
    removed_rows = []

    if filepath.lower().endswith(('.xlsx', '.xlsm')):
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active
        data = list(ws.iter_rows(values_only=True))
        if not data:
            return [], [], []
            
        headers = [str(cell) if cell is not None else "" for cell in data[0]]
        rows = data[1:]
        
        col_idx = None if col_name == "ALL COLUMNS (Check entire row)" else headers.index(col_name)

        for row in rows:
            if col_idx is None:
                should_remove = any(is_empty(v) for v in row)
            else:
                val = row[col_idx] if len(row) > col_idx else None
                should_remove = is_empty(val)

            if should_remove:
                removed_rows.append(row)
            else:
                kept_rows.append(row)
        wb.close()

    elif filepath.lower().endswith('.csv'):
        with open(filepath, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            data = list(reader)
        
        if not data:
            return [], [], []
            
        headers = data[0]
        rows = data[1:]
        
        col_idx = None if col_name == "ALL COLUMNS (Check entire row)" else headers.index(col_name)

        for row in rows:
            if col_idx is None:
                should_remove = any(is_empty(v) for v in row)
            else:
                val = row[col_idx] if len(row) > col_idx else None
                should_remove = is_empty(val)

            if should_remove:
                removed_rows.append(row)
            else:
                kept_rows.append(row)

    return headers, kept_rows, removed_rows

def save_cleaned_data(filepath, headers, kept_rows):
    """Saves the cleaned data to a new Excel file with auto-adjusted column widths."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Enriched Attendance"

    # Write Headers
    for col_idx, hdr in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=hdr)

    # Write Kept Rows
    for row_idx, row in enumerate(kept_rows, start=2):
        for col_idx, val in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)

    # Adjust Column Widths
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(headers[col_idx-1])), 10)
        for row_idx in range(2, len(kept_rows)+2):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

    wb.save(filepath)