import re
from openpyxl import load_workbook

def clean_phone(phone):
    """Remove spaces, dashes, parentheses; keep digits and commas."""
    if not phone:
        return ""
    return re.sub(r'[^\d,]', '', str(phone).strip())

def parse_phone_list(phone_str):
    """Split phone string by comma, clean each."""
    if not phone_str:
        return []
    cleaned = clean_phone(phone_str)
    parts = [p.strip() for p in cleaned.split(',') if p.strip()]
    return parts

def parse_sponsorship(filepath):
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active

    id_to_source = {}
    email_to_source = {}
    phone_to_source = {}
    email_to_id = {}
    phone_to_id = {}

    header_row = None
    col_map = {}
    
    for row in ws.iter_rows(values_only=True):
        if all(c is None for c in row):
            continue
        row_str = [str(c).strip() if c is not None else "" for c in row]
        found = {}
        for idx, val in enumerate(row_str):
            val_lower = val.lower()
            if "idnumber" in val_lower or "id" in val_lower:
                found["idnumber"] = idx
            elif "source" in val_lower or "sponsor" in val_lower:
                found["source"] = idx
            elif "email" in val_lower:
                found["email"] = idx
            elif "phone" in val_lower:
                found["phone"] = idx
            elif "full name" in val_lower or "name" in val_lower:
                found["name"] = idx
        if "idnumber" in found and "source" in found:
            col_map = found
            header_row = row
            break

    if not col_map:
        raise ValueError("Could not find required columns (IDNumber and Source) in sponsorship file.")

    for row in ws.iter_rows(values_only=True):
        if row == header_row:
            continue
        if all(c is None for c in row):
            continue
        if len(row) <= max(col_map.values()):
            continue

        idnumber = str(row[col_map["idnumber"]]).strip() if row[col_map["idnumber"]] else ""
        source = str(row[col_map["source"]]).strip() if row[col_map["source"]] else ""
        email = str(row[col_map.get("email", 0)]).strip() if "email" in col_map and row[col_map["email"]] else ""
        phone_raw = str(row[col_map.get("phone", 0)]).strip() if "phone" in col_map and row[col_map["phone"]] else ""

        if idnumber and source:
            id_to_source[idnumber] = source
            if email:
                email_key = email.lower()
                email_to_source[email_key] = source
                email_to_id[email_key] = idnumber
            if phone_raw:
                phone_list = parse_phone_list(phone_raw)
                for p in phone_list:
                    if p:
                        phone_to_source[p] = source
                        phone_to_id[p] = idnumber

    return id_to_source, email_to_source, phone_to_source, email_to_id, phone_to_id

def parse_attendance(filepath):
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active

    header_row = None
    col_map = {}

    for row in ws.iter_rows(values_only=True):
        if all(c is None for c in row):
            continue
        row_str = [str(c).strip() if c is not None else "" for c in row]
        found = {}
        for idx, val in enumerate(row_str):
            val_lower = val.lower()
            if "name" in val_lower:
                found["name"] = idx
            elif "email" in val_lower:
                found["email"] = idx
            elif "phone" in val_lower:
                found["phone"] = idx
            elif "idnumber" in val_lower or "id" in val_lower:
                found["idnumber"] = idx
            elif "join" in val_lower:
                found["join"] = idx
            elif "left" in val_lower:
                found["left"] = idx
            elif "duration" in val_lower:
                found["duration"] = idx
        if "name" in found and "join" in found and "left" in found and "duration" in found:
            col_map = found
            header_row = row
            break

    if not col_map:
        first_row = None
        for row in ws.iter_rows(values_only=True):
            if all(c is None for c in row):
                continue
            first_row = row
            break
        if first_row:
            row_str = [str(c).strip().lower() if c else "" for c in first_row]
            if any("name" in s for s in row_str):
                col_map = {"name": 0, "email": 1, "phone": 2, "idnumber": 3, "join": 4, "left": 5, "duration": 6}
                header_row = first_row
            else:
                col_map = {"name": 0, "email": 1, "phone": 2, "idnumber": 3, "join": 4, "left": 5, "duration": 6}
                header_row = None

    att_list = []
    for row in ws.iter_rows(values_only=True):
        if header_row is not None and row == header_row:
            continue
        if all(c is None for c in row):
            continue
        if len(row) <= max(col_map.values()):
            continue

        name = str(row[col_map["name"]]).strip() if row[col_map["name"]] else ""
        if not name:
            continue

        email = str(row[col_map.get("email", 0)]).strip() if "email" in col_map and row[col_map["email"]] else ""
        phone = str(row[col_map.get("phone", 0)]).strip() if "phone" in col_map and row[col_map["phone"]] else ""
        idnumber = str(row[col_map.get("idnumber", 0)]).strip() if "idnumber" in col_map and row[col_map["idnumber"]] else ""
        join = row[col_map["join"]]
        left = row[col_map["left"]]
        dur = row[col_map["duration"]]

        if isinstance(dur, (int, float)):
            h = int(dur // 3600)
            m = int((dur % 3600) // 60)
            s = int(dur % 60)
            dur_str = f"{h:02d}:{m:02d}:{s:02d}"
        else:
            dur_str = str(dur).strip()

        att_list.append({
            "name": name, "email": email, "phone": phone, "idnumber": idnumber,
            "join": join, "left": left, "duration": dur_str
        })

    return att_list

def match_sources(sponsor_path, att_path):
    """Main logic to enrich attendance with source data."""
    id_to_source, email_to_source, phone_to_source, email_to_id, phone_to_id = parse_sponsorship(sponsor_path)
    attendance = parse_attendance(att_path)
    
    if not attendance:
        raise ValueError("No attendance data found.")

    matched_count = 0
    unmatched_names = []
    result_data = []

    for att in attendance:
        name = att["name"]
        email = att.get("email", "").strip()
        phone = att.get("phone", "").strip()
        idnumber = att.get("idnumber", "").strip()
        source = ""

        if idnumber:
            if idnumber in id_to_source:
                source = id_to_source[idnumber]
                matched_count += 1
            else:
                if email and email.lower() in email_to_source:
                    source = email_to_source[email.lower()]
                    matched_count += 1
                elif phone and clean_phone(phone) in phone_to_source:
                    source = phone_to_source[clean_phone(phone)]
                    matched_count += 1
                else:
                    unmatched_names.append(name)
        else:
            if email and email.lower() in email_to_id:
                idnumber = email_to_id[email.lower()]
                source = email_to_source[email.lower()]
                matched_count += 1
            elif phone and clean_phone(phone) in phone_to_id:
                idnumber = phone_to_id[clean_phone(phone)]
                source = phone_to_source[clean_phone(phone)]
                matched_count += 1
            else:
                unmatched_names.append(name)

        row = {
            "Name": name, "Email": email, "Phone": phone, "Source": source,
            "IDNumber": idnumber, "Join": att["join"], "Left": att["left"], "Duration": att["duration"]
        }
        result_data.append(row)

    return result_data, unmatched_names, matched_count, len(attendance)