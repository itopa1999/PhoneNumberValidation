import re
import difflib
from openpyxl import load_workbook

def clean_name(name):
    """Normalize name: lower case, remove punctuation, sort tokens."""
    if not name:
        return ""
    s = str(name).strip()
    s = re.sub(r'[^\w\s]', '', s)
    tokens = s.lower().split()
    tokens = [t for t in tokens if t]
    return " ".join(sorted(tokens))

def parse_duration(value):
    if value is None:
        return 0
    s = str(value).strip()
    parts = s.split(':')
    if len(parts) == 3:
        h, m, sec = map(int, parts)
        return h * 3600 + m * 60 + sec
    elif len(parts) == 2:
        m, sec = map(int, parts)
        return m * 60 + sec
    else:
        try:
            return int(s)
        except:
            return 0

def format_duration(seconds):
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    return f"{h:02d}:{m:02d}:{s:02d}"

def parse_registration(filepath):
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    data = []
    header_row = None
    col_map = {}

    for row in ws.iter_rows(values_only=True):
        if all(c is None for c in row):
            continue
        row_str = [str(c).strip() if c is not None else "" for c in row]
        found = {}
        for idx, val in enumerate(row_str):
            val_lower = val.lower()
            if "email" in val_lower:
                found["email"] = idx
            elif "name" in val_lower:
                found["name"] = idx
            elif "idnumber" in val_lower or "id" in val_lower:
                found["idnumber"] = idx
            elif "phone" in val_lower:
                found["phone"] = idx
        if "email" in found and "name" in found and "idnumber" in found:
            col_map = found
            header_row = row
            break

    if not col_map:
        first_row = None
        for row in ws.iter_rows(values_only=True):
            if all(c is None for c in row):
                continue
            first_row = row
            break
        if first_row and len(first_row) >= 4:
            email_col = None
            name_col = None
            id_col = None
            phone_col = None
            for idx, val in enumerate(first_row):
                if val:
                    val_lower = str(val).lower()
                    if '@' in val_lower:
                        email_col = idx
                    elif 'id' in val_lower:
                        id_col = idx
                    elif 'phone' in val_lower:
                        phone_col = idx
            if email_col is None: email_col = 1
            if id_col is None: id_col = 3
            if phone_col is None: phone_col = 4 if len(first_row) > 4 else None
            
            name_col = None
            for idx in range(len(first_row)):
                if idx not in (email_col, id_col, phone_col) if phone_col is not None else (email_col, id_col):
                    name_col = idx
                    break
            if name_col is None: name_col = 2
            
            col_map = {"email": email_col, "name": name_col, "idnumber": id_col}
            if phone_col is not None:
                col_map["phone"] = phone_col
            header_row = first_row
        else:
            return []

    for row in ws.iter_rows(values_only=True):
        if row == header_row:
            continue
        if all(c is None for c in row):
            continue
        if len(row) <= max(col_map.values()):
            continue
        name = str(row[col_map["name"]]).strip() if row[col_map["name"]] else ""
        email = str(row[col_map["email"]]).strip() if row[col_map["email"]] else ""
        idnumber = str(row[col_map["idnumber"]]).strip() if row[col_map["idnumber"]] else ""
        phone = ""
        if "phone" in col_map:
            phone = str(row[col_map["phone"]]).strip() if row[col_map["phone"]] else ""
        if name:
            data.append({"name": name, "email": email, "idnumber": idnumber, "phone": phone})
    return data

def parse_attendance(filepath):
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    att_list = []
    first_row = None
    for row in ws.iter_rows(values_only=True):
        if all(c is None for c in row):
            continue
        first_row = row
        break
    if first_row is None:
        return []

    header = False
    if first_row and len(first_row) >= 4:
        if any("name" in str(c).lower() for c in first_row if c):
            header = True
        elif any("join" in str(c).lower() for c in first_row if c):
            header = True

    start_row = 1 if header else 0
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx < start_row:
            continue
        if all(c is None for c in row):
            continue
        if len(row) < 4:
            continue
        name = str(row[0]).strip() if row[0] else ""
        if not name:
            continue
        join = row[1]
        left = row[2]
        dur = row[3]
        dur_sec = parse_duration(dur)
        att_list.append({
            "name": name,
            "join": join,
            "left": left,
            "duration_seconds": dur_sec
        })
    return att_list

def merge_data(reg_path, att_path):
    """Main logic to merge attendance and registration data."""
    reg_data = parse_registration(reg_path)
    if not reg_data:
        raise ValueError("No registration data found or file format is invalid.")

    reg_map = {}
    for item in reg_data:
        cleaned = clean_name(item["name"])
        if cleaned not in reg_map:
            reg_map[cleaned] = (item["email"], item["idnumber"], item.get("phone", ""))

    attendance = parse_attendance(att_path)
    if not attendance:
        raise ValueError("No attendance data found or file format is invalid.")

    matched_count = 0
    unmatched_names = []
    result_data = []

    for att in attendance:
        name = att["name"]
        cleaned = clean_name(name)
        email, idnumber, phone = "", "", ""

        if cleaned in reg_map:
            email, idnumber, phone = reg_map[cleaned]
            matched_count += 1
        else:
            candidates = difflib.get_close_matches(cleaned, reg_map.keys(), n=1, cutoff=0.7)
            if candidates:
                email, idnumber, phone = reg_map[candidates[0]]
                matched_count += 1
            else:
                unmatched_names.append(name)

        row = {
            "Name": name,
            "Email": email,
            "Phone": phone,
            "IDNumber": idnumber,
            "Join": att["join"],
            "Left": att["left"],
            "Duration": format_duration(att["duration_seconds"])
        }
        result_data.append(row)

    return result_data, unmatched_names, matched_count, len(attendance)