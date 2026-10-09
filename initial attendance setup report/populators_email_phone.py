import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
import re
import os

class EmailPhoneAdderApp:
    def __init__(self, root):
        self.root = root
        self.root.title("📧 Add Email & Phone – Choose Columns")
        self.root.geometry("1000x820")
        self.root.configure(bg="#f0f4f8")

        self.master_path = tk.StringVar()
        self.target_path = tk.StringVar()

        # Column selection variables
        self.name_col = tk.StringVar()          # single "Name" (optional if first/last used)
        self.first_name_col = tk.StringVar()    # <-- NEW
        self.last_name_col = tk.StringVar()     # <-- NEW
        self.email_col = tk.StringVar()
        self.phone_col = tk.StringVar()
        self.id_col = tk.StringVar()
        self.target_headers = []  # list of header strings

        self.result_data = []
        self.unmatched_names = []

        self.setup_styles()

        main = ttk.Frame(root, padding="15")
        main.pack(fill="both", expand=True)

        # Title
        ttk.Label(main, text="📧 Add Email & Phone – Custom Column Selection",
                  font=("Segoe UI", 16, "bold"), foreground="#1a2b3c").pack(anchor="w")
        ttk.Label(main, text="Select which columns in the target file contain Name, Email, and Phone.",
                  font=("Segoe UI", 10), foreground="#5a6b7c").pack(anchor="w", pady=(0, 10))

        ttk.Separator(main, orient="horizontal").pack(fill="x", pady=(0, 10))

        # File selection
        master_frame = ttk.Frame(main)
        master_frame.pack(fill="x", pady=3)
        ttk.Label(master_frame, text="📁 Master File (with Email & Phone):").pack(side="left", padx=(0, 8))
        ttk.Entry(master_frame, textvariable=self.master_path, width=60).pack(side="left", fill="x", expand=True, padx=(0, 8))
        ttk.Button(master_frame, text="Browse", command=lambda: self.browse_file(self.master_path)).pack(side="left")

        target_frame = ttk.Frame(main)
        target_frame.pack(fill="x", pady=3)
        ttk.Label(target_frame, text="📁 Target File (to enrich):").pack(side="left", padx=(0, 8))
        ttk.Entry(target_frame, textvariable=self.target_path, width=60).pack(side="left", fill="x", expand=True, padx=(0, 8))
        ttk.Button(target_frame, text="Browse", command=lambda: self.browse_target()).pack(side="left")

        # Column selection frame
        col_frame = ttk.LabelFrame(main, text="Target Column Mapping", padding=10)
        col_frame.pack(fill="x", pady=10)

        # ID column
        ttk.Label(col_frame, text="ID column (optional – used first for matching):").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        self.id_combo = ttk.Combobox(col_frame, textvariable=self.id_col, state="readonly", width=28)
        self.id_combo.grid(row=0, column=1, padx=5, pady=5, sticky="w")

        # Name column (single)
        ttk.Label(col_frame, text="Name column (optional if using First/Last):").grid(row=0, column=2, padx=5, pady=5, sticky="w")
        self.name_combo = ttk.Combobox(col_frame, textvariable=self.name_col, state="readonly", width=28)
        self.name_combo.grid(row=0, column=3, padx=5, pady=5, sticky="w")

        # First name column  <-- NEW
        ttk.Label(col_frame, text="First Name column (optional):").grid(row=1, column=0, padx=5, pady=5, sticky="w")
        self.first_combo = ttk.Combobox(col_frame, textvariable=self.first_name_col, state="readonly", width=28)
        self.first_combo.grid(row=1, column=1, padx=5, pady=5, sticky="w")

        # Last name column  <-- NEW
        ttk.Label(col_frame, text="Last Name column (optional):").grid(row=1, column=2, padx=5, pady=5, sticky="w")
        self.last_combo = ttk.Combobox(col_frame, textvariable=self.last_name_col, state="readonly", width=28)
        self.last_combo.grid(row=1, column=3, padx=5, pady=5, sticky="w")

        # Email column
        ttk.Label(col_frame, text="Email column (optional):").grid(row=2, column=0, padx=5, pady=5, sticky="w")
        self.email_combo = ttk.Combobox(col_frame, textvariable=self.email_col, state="readonly", width=28)
        self.email_combo.grid(row=2, column=1, padx=5, pady=5, sticky="w")

        # Phone column
        ttk.Label(col_frame, text="Phone column (optional):").grid(row=2, column=2, padx=5, pady=5, sticky="w")
        self.phone_combo = ttk.Combobox(col_frame, textvariable=self.phone_col, state="readonly", width=28)
        self.phone_combo.grid(row=2, column=3, padx=5, pady=5, sticky="w")

        # Refresh button
        ttk.Button(col_frame, text="🔄 Refresh Columns", command=self.load_target_headers).grid(row=3, column=3, padx=5, pady=5)

        # Buttons
        btn_frame = ttk.Frame(main)
        btn_frame.pack(fill="x", pady=10)

        ttk.Button(btn_frame, text="⚡ Process & Enrich", command=self.process_files,
                   style="Success.TButton").pack(side="left", padx=(0, 10))
        ttk.Button(btn_frame, text="💾 Save Enriched File", command=self.save_result,
                   style="Primary.TButton").pack(side="left", padx=(0, 10))
        ttk.Button(btn_frame, text="Clear All", command=self.clear_all,
                   style="Danger.TButton").pack(side="left")

        self.status = ttk.Label(main, text="Ready", relief="sunken", anchor="w", style="Status.TLabel")
        self.status.pack(fill="x", pady=(5, 10))

        ttk.Label(main, text="Unmatched Records:", font=("Segoe UI", 10, "bold")).pack(anchor="w")
        self.unmatched_text = scrolledtext.ScrolledText(main, height=8, font=("Segoe UI", 9), state="disabled")
        self.unmatched_text.pack(fill="both", expand=True, pady=5)

        self.count_label = ttk.Label(main, text="", font=("Segoe UI", 9, "italic"))
        self.count_label.pack(anchor="w", pady=(0, 5))

        self.setup_styles()

    def setup_styles(self):
        style = ttk.Style()
        style.theme_use("clam")
        style.configure("TFrame", background="#f0f4f8")
        style.configure("TLabel", background="#f0f4f8", foreground="#1a2b3c", font=("Segoe UI", 10))
        style.configure("TButton", font=("Segoe UI", 10, "bold"), padding=6)
        style.map("TButton", background=[("active", "#2a7de1")])
        style.configure("Success.TButton", background="#28a745", foreground="white")
        style.map("Success.TButton", background=[("active", "#218838")])
        style.configure("Primary.TButton", background="#2a7de1", foreground="white")
        style.map("Primary.TButton", background=[("active", "#1a5fb4")])
        style.configure("Danger.TButton", background="#dc3545", foreground="white")
        style.map("Danger.TButton", background=[("active", "#c82333")])
        style.configure("Status.TLabel", background="#e9ecef", foreground="#1a2b3c", font=("Segoe UI", 9), padding=4)

    def browse_file(self, var):
        f = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")])
        if f:
            var.set(f)

    def browse_target(self):
        f = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")])
        if f:
            self.target_path.set(f)
            self.load_target_headers()

    def load_target_headers(self):
        """Read target file headers and populate comboboxes."""
        filepath = self.target_path.get()
        if not filepath or not os.path.exists(filepath):
            return
        try:
            wb = load_workbook(filepath, data_only=True)
            ws = wb.active
            # find first non-empty row as header
            for row in ws.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    self.target_headers = [str(c).strip() if c is not None else "" for c in row]
                    break
            else:
                self.target_headers = []
            if not self.target_headers:
                # fallback: use column letters
                max_col = ws.max_column
                self.target_headers = [f"Column {get_column_letter(i+1)}" for i in range(max_col)]
            # Set combobox values
            self.id_combo['values'] = self.target_headers
            self.name_combo['values'] = self.target_headers
            self.first_combo['values'] = self.target_headers
            self.last_combo['values'] = self.target_headers
            self.email_combo['values'] = self.target_headers
            self.phone_combo['values'] = self.target_headers
            # Try auto-detect by common keywords
            auto_id = None
            auto_name = None
            auto_first = None
            auto_last = None
            auto_email = None
            auto_phone = None
            for idx, h in enumerate(self.target_headers):
                h_lower = h.lower().strip()
                # Normalize separators for detection
                h_norm = h_lower.replace("_", " ").replace("-", " ")
                if auto_id is None and ("id" in h_norm or "number" in h_norm or h_norm == "no"):
                    auto_id = h
                if auto_name is None and "name" in h_norm and "first" not in h_norm and "last" not in h_norm and "surname" not in h_norm:
                    auto_name = h
                if auto_first is None and ("first" in h_norm and "name" in h_norm or h_norm in ("firstname", "first name", "given name", "givenname")):
                    auto_first = h
                if auto_last is None and ("last" in h_norm and "name" in h_norm or "surname" in h_norm or h_norm in ("lastname", "last name", "family name", "familyname")):
                    auto_last = h
                if auto_email is None and "email" in h_norm:
                    auto_email = h
                if auto_phone is None and "phone" in h_norm:
                    auto_phone = h
            if auto_id:
                self.id_col.set(auto_id)
            if auto_first:
                self.first_name_col.set(auto_first)
            if auto_last:
                self.last_name_col.set(auto_last)
            if auto_name:
                self.name_col.set(auto_name)
            elif not auto_first and not auto_last and self.target_headers:
                self.name_col.set(self.target_headers[0])
            if auto_email:
                self.email_col.set(auto_email)
            if auto_phone:
                self.phone_col.set(auto_phone)
            self.status.config(text="Target columns loaded. Select the correct ones if needed.")
        except Exception as e:
            messagebox.showerror("Error", f"Could not read target file:\n{str(e)}")

    # ---------- Matching helpers ----------
    def clean_name_tokens(self, name):
        if not name:
            return set()
        name = re.sub(r'\b(mr|mrs|ms|dr|prof|engr|hon|rev|sir|madam|chief)\b', '', name, flags=re.IGNORECASE)
        name = re.sub(r'[^a-zA-Z0-9\s]', '', name)
        words = set(name.lower().split())
        words = {w for w in words if len(w) > 1}
        return words

    def name_match_score(self, target_tokens, master_tokens):
        common = target_tokens & master_tokens
        score = len(common)
        min_len = min(len(target_tokens), len(master_tokens))
        if min_len == 0:
            return score, False
        coverage = score / min_len
        return score, (score >= 2 or coverage >= 0.6)

    def normalize_id(self, value):
        """Normalize an ID value to a comparable string (strip spaces, lowercase)."""
        if value is None:
            return ""
        return str(value).strip().lower()

    def build_full_name(self, first, last):
        """Combine first and last names into a single string."""
        parts = []
        if first:
            parts.append(str(first).strip())
        if last:
            parts.append(str(last).strip())
        return " ".join(p for p in parts if p)

    # ---------- Process ----------
    def process_files(self):
        if not self.master_path.get() or not self.target_path.get():
            messagebox.showerror("Error", "Please select both files.")
            return

        # Validate column selection: need either Name OR First/Last
        if not self.name_col.get() and not self.first_name_col.get() and not self.last_name_col.get():
            messagebox.showerror("Error", "Please select either a Name column, or First Name / Last Name columns.")
            return

        try:
            self.status.config(text="Reading master file...")
            master_data, master_headers = self.read_master(self.master_path.get())
            if not master_data:
                messagebox.showerror("Error", "No valid master data found (need Name or First/Last, Email, Phone).")
                return

            self.status.config(text="Reading target file...")
            target_rows, target_headers = self.read_target_with_mapping(
                self.target_path.get(),
                self.name_col.get() if self.name_col.get() else None,
                self.email_col.get() if self.email_col.get() else None,
                self.phone_col.get() if self.phone_col.get() else None,
                self.id_col.get() if self.id_col.get() else None,
                self.first_name_col.get() if self.first_name_col.get() else None,
                self.last_name_col.get() if self.last_name_col.get() else None
            )
            if not target_rows:
                messagebox.showerror("Error", "No data in target file.")
                return

            # Build lookup dictionaries
            id_to_record = {}
            email_to_record = {}
            phone_to_record = {}
            master_token_list = []

            for master in master_data:
                name = master["name"].strip()
                email = master["email"].strip().lower() if master["email"] else ""
                phone = master["phone"].strip() if master["phone"] else ""
                id_num = self.normalize_id(master.get("id", ""))

                if id_num:
                    id_to_record[id_num] = master
                if email:
                    email_to_record[email] = master
                if phone:
                    phone_to_record[phone] = master
                if name:
                    tokens = self.clean_name_tokens(name)
                    if tokens:
                        master_token_list.append((master, tokens))

            # Process target rows
            self.unmatched_names = []
            enriched_rows = []

            for row in target_rows:
                name = row.get("name", "").strip()
                email = row.get("email", "").strip().lower() if "email" in row else ""
                phone = row.get("phone", "").strip() if "phone" in row else ""
                id_num = self.normalize_id(row.get("id", ""))

                matched_email = ""
                matched_phone = ""
                matched = False

                # 0) Exact ID match (highest priority)
                if id_num and id_num in id_to_record:
                    rec = id_to_record[id_num]
                    matched_email = rec["email"]
                    matched_phone = rec["phone"]
                    matched = True

                # 1) Exact email
                if not matched and email and email in email_to_record:
                    rec = email_to_record[email]
                    matched_email = rec["email"]
                    matched_phone = rec["phone"]
                    matched = True

                # 2) Exact phone
                if not matched and phone and phone in phone_to_record:
                    rec = phone_to_record[phone]
                    matched_email = rec["email"]
                    matched_phone = rec["phone"]
                    matched = True

                # 3) Name keyword matching
                if not matched and name:
                    target_tokens = self.clean_name_tokens(name)
                    best_score = 0
                    best_rec = None
                    for m_rec, m_tokens in master_token_list:
                        score, good = self.name_match_score(target_tokens, m_tokens)
                        if good and score > best_score:
                            best_score = score
                            best_rec = m_rec
                    if best_rec:
                        matched_email = best_rec["email"]
                        matched_phone = best_rec["phone"]
                        matched = True

                if not matched:
                    self.unmatched_names.append(name if name else "(no name)")

                new_row = row.copy()
                new_row["added_email"] = matched_email
                new_row["added_phone"] = matched_phone
                enriched_rows.append(new_row)

            self.result_data = enriched_rows
            self.display_unmatched()

            total = len(target_rows)
            matched = total - len(self.unmatched_names)
            self.status.config(text=f"✅ Matching complete: {matched}/{total} matched.")
            self.count_label.config(text=f"{total} target rows, {matched} matched, {len(self.unmatched_names)} unmatched.")

            if self.unmatched_names:
                messagebox.showwarning("Unmatched Records",
                                       f"{len(self.unmatched_names)} records could not be matched.\n"
                                       "Check the Unmatched list below.")

        except Exception as e:
            messagebox.showerror("Error", f"Processing failed:\n{str(e)}")
            self.status.config(text="❌ Error")

    # ---------- File reading with custom mapping ----------
    def read_master(self, filepath):
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active
        header_row = None
        col_map = {}
        for row in ws.iter_rows(values_only=True):
            if all(c is None for c in row):
                continue
            row_str = [str(c).strip().lower() if c is not None else "" for c in row]
            found = {}
            for idx, val in enumerate(row_str):
                # Normalize
                v = val.replace("_", " ").replace("-", " ")
                if found.get("name") is None and "name" in v and "first" not in v and "last" not in v and "surname" not in v:
                    found["name"] = idx
                elif found.get("first") is None and (("first" in v and "name" in v) or v in ("firstname", "first name", "given name", "givenname")):
                    found["first"] = idx
                elif found.get("last") is None and (("last" in v and "name" in v) or "surname" in v or v in ("lastname", "last name", "family name", "familyname")):
                    found["last"] = idx
                elif found.get("email") is None and "email" in v:
                    found["email"] = idx
                elif found.get("phone") is None and "phone" in v:
                    found["phone"] = idx
                elif found.get("id") is None and ("id" in v or "number" in v or v == "no"):
                    found["id"] = idx

            # Need email & phone, plus either name or first/last
            has_name = "name" in found or ("first" in found and "last" in found) or "first" in found or "last" in found
            if has_name and "email" in found and "phone" in found:
                col_map = found
                header_row = row
                break

        if not col_map:
            messagebox.showerror("Error", "Master file must have columns: Name (or First/Last), Email, Phone.")
            return None, None

        def get_val(row, key):
            if key in col_map and len(row) > col_map[key] and row[col_map[key]] is not None:
                return str(row[col_map[key]]).strip()
            return ""

        data = []
        for row in ws.iter_rows(values_only=True):
            if header_row is not None and row == header_row:
                continue
            if all(c is None for c in row):
                continue
            if len(row) <= max(col_map.values()):
                continue

            name = get_val(row, "name")
            first = get_val(row, "first")
            last = get_val(row, "last")
            # If no single name, build from first/last
            if not name:
                name = self.build_full_name(first, last)

            email = get_val(row, "email")
            phone = get_val(row, "phone")
            id_val = get_val(row, "id")

            if name:
                data.append({"name": name, "email": email, "phone": phone, "id": id_val})
        return data, header_row

    def read_target_with_mapping(self, filepath, name_col, email_col, phone_col, id_col=None,
                                 first_name_col=None, last_name_col=None):
        """Read target file using user-specified column headers."""
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active

        # Find header row
        header_row = None
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                header_row = [str(c).strip() if c is not None else "" for c in row]
                break
        if not header_row:
            return None, None

        # Find indices of selected columns
        name_idx = header_row.index(name_col) if name_col and name_col in header_row else None
        first_idx = header_row.index(first_name_col) if first_name_col and first_name_col in header_row else None
        last_idx = header_row.index(last_name_col) if last_name_col and last_name_col in header_row else None
        email_idx = header_row.index(email_col) if email_col and email_col in header_row else None
        phone_idx = header_row.index(phone_col) if phone_col and phone_col in header_row else None
        id_idx = header_row.index(id_col) if id_col and id_col in header_row else None

        if name_idx is None and first_idx is None and last_idx is None:
            raise ValueError("Please select a Name column, or a First Name / Last Name column.")

        # Max index we need to look at
        max_needed = max([i for i in [name_idx, first_idx, last_idx, email_idx, phone_idx, id_idx] if i is not None])

        data = []
        for row in ws.iter_rows(values_only=True):
            if row == header_row or all(c is None for c in row):
                continue
            if len(row) <= max_needed:
                continue
            row_dict = {}
            # store all columns by header
            for idx, val in enumerate(row):
                if idx < len(header_row):
                    hdr = header_row[idx]
                    row_dict[hdr] = val if val is not None else ""
            # Build name value: prefer single Name, else combine first + last
            name_val = ""
            if name_idx is not None and row[name_idx]:
                name_val = str(row[name_idx]).strip()
            if not name_val:
                first = str(row[first_idx]).strip() if first_idx is not None and row[first_idx] else ""
                last = str(row[last_idx]).strip() if last_idx is not None and row[last_idx] else ""
                name_val = self.build_full_name(first, last)
            if not name_val:
                continue  # skip rows with no name at all

            row_dict["name"] = name_val
            if email_idx is not None:
                row_dict["email"] = str(row[email_idx]).strip() if row[email_idx] else ""
            else:
                row_dict["email"] = ""
            if phone_idx is not None:
                row_dict["phone"] = str(row[phone_idx]).strip() if row[phone_idx] else ""
            else:
                row_dict["phone"] = ""
            if id_idx is not None:
                row_dict["id"] = str(row[id_idx]).strip() if row[id_idx] else ""
            else:
                row_dict["id"] = ""
            data.append(row_dict)
        return data, header_row

    def display_unmatched(self):
        self.unmatched_text.config(state="normal")
        self.unmatched_text.delete("1.0", tk.END)
        if self.unmatched_names:
            for name in self.unmatched_names:
                self.unmatched_text.insert(tk.END, name + "\n")
        self.unmatched_text.config(state="disabled")

    def save_result(self):
        if not self.result_data:
            messagebox.showwarning("Warning", "No data to save. Process files first.")
            return

        file_path = filedialog.asksaveasfilename(
            title="Save Enriched File",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        if not file_path:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Enriched"

            if self.result_data:
                first_row = self.result_data[0]
                # Preserve the original target columns (by header name), in original order
                original_keys = [k for k in first_row.keys()
                                 if k not in ('added_email', 'added_phone', 'name', 'email', 'phone', 'id')]
                headers = original_keys + ['Name', 'Email', 'Phone']
            else:
                headers = ['Name', 'Email', 'Phone']

            for col_idx, hdr in enumerate(headers, 1):
                ws.cell(row=1, column=col_idx, value=hdr)

            for row_idx, row_data in enumerate(self.result_data, start=2):
                for col_idx, hdr in enumerate(headers, 1):
                    if hdr == 'Name':
                        val = row_data.get('name', '')
                    elif hdr == 'Email':
                        val = row_data.get('added_email', '')
                    elif hdr == 'Phone':
                        val = row_data.get('added_phone', '')
                    else:
                        val = row_data.get(hdr, '')
                    ws.cell(row=row_idx, column=col_idx, value=val)

            for col_idx in range(1, len(headers)+1):
                col_letter = get_column_letter(col_idx)
                max_len = max(len(str(headers[col_idx-1])), 10)
                for row_idx in range(2, len(self.result_data)+2):
                    val = ws.cell(row=row_idx, column=col_idx).value
                    if val:
                        max_len = max(max_len, len(str(val)))
                ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

            wb.save(file_path)
            messagebox.showinfo("Success", f"File saved to:\n{file_path}")
            self.status.config(text=f"💾 Saved: {os.path.basename(file_path)}")
        except Exception as e:
            messagebox.showerror("Error", f"Could not save file:\n{str(e)}")

    def clear_all(self):
        self.master_path.set("")
        self.target_path.set("")
        self.id_col.set("")
        self.name_col.set("")
        self.first_name_col.set("")
        self.last_name_col.set("")
        self.email_col.set("")
        self.phone_col.set("")
        self.id_combo['values'] = []
        self.name_combo['values'] = []
        self.first_combo['values'] = []
        self.last_combo['values'] = []
        self.email_combo['values'] = []
        self.phone_combo['values'] = []
        self.result_data = []
        self.unmatched_names = []
        self.unmatched_text.config(state="normal")
        self.unmatched_text.delete("1.0", tk.END)
        self.unmatched_text.config(state="disabled")
        self.status.config(text="Cleared")
        self.count_label.config(text="")

if __name__ == "__main__":
    root = tk.Tk()
    app = EmailPhoneAdderApp(root)
    root.mainloop()