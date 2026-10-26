import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import os
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# Logic files
from logic.coordinator_remover import remove_coordinators, get_file_headers
from logic.duplicate_name import process_duplicates
from logic.email_phone_merger import merge_data
from logic.data_cleaner import get_file_headers as get_headers_dc, process_file_data
from logic.source_adder import match_sources


class MasterDashboardApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Data & Attendance Master Suite")
        self.root.geometry("1150x750")

        # --- Colors ---
        self.bg_toolbar = "#0f172a"
        self.bg_sidebar = "#1e293b"
        self.bg_header = "#ffffff"
        self.bg_main = "#f0f4f8"
        self.accent_color = "#2a7de1"
        self.text_dark = "#1a2b3c"
        self.text_light = "#64748b"
        self.border_color = "#e2e8f0"

        self.setup_styles()
        self.build_layout()
        self.show_dashboard_home()

    def setup_styles(self):
        style = ttk.Style()
        style.theme_use('clam')

        # Sidebar
        style.configure("SidebarTitle.TLabel", background=self.bg_sidebar, foreground="#ffffff",
                        font=("Segoe UI", 16, "bold"))
        style.configure("SidebarFolder.TButton", background=self.bg_sidebar, foreground="#f8fafc",
                        font=("Segoe UI", 11, "bold"), borderwidth=0, padding=10, anchor="w")
        style.map("SidebarFolder.TButton", background=[("active", "#334155")])
        style.configure("SidebarItem.TButton", background=self.bg_sidebar, foreground="#94a3b8",
                        font=("Segoe UI", 10), borderwidth=0, padding=8, anchor="w")
        style.map("SidebarItem.TButton", background=[("active", "#334155")],
                  foreground=[("active", "#ffffff")])

        # Launch
        style.configure("Launch.TButton", background=self.accent_color, foreground="white",
                        font=("Segoe UI", 10, "bold"), padding=8)
        style.map("Launch.TButton", background=[("active", "#1a5fb4")])

        # Tool buttons
        style.configure("Success.TButton", background="#28a745", foreground="white")
        style.map("Success.TButton", background=[("active", "#218838")])
        style.configure("Primary.TButton", background=self.accent_color, foreground="white")
        style.map("Primary.TButton", background=[("active", "#1a5fb4")])
        style.configure("Danger.TButton", background="#dc3545", foreground="white")
        style.map("Danger.TButton", background=[("active", "#c82333")])
        style.configure("Status.TLabel", background="#e9ecef", foreground=self.text_dark,
                        font=("Segoe UI", 9), padding=4)

        style.configure("Treeview", background="white", foreground=self.text_dark,
                        fieldbackground="white", font=("Segoe UI", 9))
        style.map("Treeview", background=[("selected", self.accent_color)])
        style.configure("Treeview.Heading", background="#d9e2ec", foreground=self.text_dark,
                        font=("Segoe UI", 10, "bold"))

    def build_layout(self):
        # TOOLBAR
        toolbar = tk.Frame(self.root, bg=self.bg_toolbar, height=60)
        toolbar.pack(side="top", fill="x")
        toolbar.pack_propagate(False)
        tk.Label(toolbar, text="⚡ DATA SUITE", bg=self.bg_toolbar, fg="white",
                 font=("Segoe UI", 14, "bold"), padx=20).pack(side="left")
        tk.Label(toolbar, text="Admin User 👤", bg=self.bg_toolbar, fg="#94a3b8",
                 font=("Segoe UI", 10), padx=20).pack(side="right")

        workspace = tk.Frame(self.root, bg=self.bg_main)
        workspace.pack(side="top", fill="both", expand=True)

        # SIDEBAR
        sidebar = tk.Frame(workspace, bg=self.bg_sidebar, width=250)
        sidebar.pack(side="left", fill="y")
        sidebar.pack_propagate(False)

        title_frame = tk.Frame(sidebar, bg=self.bg_sidebar, pady=15)
        title_frame.pack(fill="x")
        ttk.Label(title_frame, text="App", style="SidebarTitle.TLabel", anchor="w",
                  padding=(20, 0)).pack(fill="x")
        tk.Frame(sidebar, bg="#334155", height=1).pack(fill="x", padx=20, pady=(0, 10))

        self.create_collapsible_folder(sidebar, "Attendance setup report", [
            ("Remove Duplicate name", self.load_deduplicator),
            ("Add Email and Phone", self.load_email_phone_merger),
            ("Remove Coordinators", self.load_remove_coordinators),
            ("Add Particpaints Source", self.load_source_adder),
        ])
        self.create_collapsible_folder(sidebar, "Initial attendance setup report", [
            ("Populate Email and Phone", lambda: self.load_placeholder("Populate Email and Phone")),
            ("Revalidate Participants Source", lambda: self.load_placeholder("Revalidate Participants")),
        ])
        self.create_collapsible_folder(sidebar, "Send LMS login", [
            ("Email LMS Credential Links Generator", lambda: self.load_placeholder("Email LMS Generator")),
            ("Whatsapp LMS Credential Links Generator", lambda: self.load_placeholder("WhatsApp LMS Generator")),
        ])
        self.create_collapsible_folder(sidebar, "Utils", [
            ("Data Cleaner", self.load_data_cleaner),
            ("Generate Auto Report", lambda: self.load_placeholder("Auto Report")),
        ])

        tk.Label(sidebar, text="v1.0.0 • Stable", bg=self.bg_sidebar, fg="#64748b",
                 font=("Segoe UI", 9), anchor="w", padx=20).pack(side="bottom", pady=20)

        # RIGHT SIDE
        self.right_container = tk.Frame(workspace, bg=self.bg_main)
        self.right_container.pack(side="right", fill="both", expand=True)

        self.header_frame = tk.Frame(self.right_container, bg=self.bg_header, height=70,
                                     highlightbackground=self.border_color, highlightthickness=1)
        self.header_frame.pack(side="top", fill="x")
        self.header_frame.pack_propagate(False)

        self.header_title = tk.Label(self.header_frame, text="Dashboard Overview",
                                     font=("Segoe UI", 16, "bold"), bg=self.bg_header,
                                     fg=self.text_dark, padx=30)
        self.header_title.pack(side="left", pady=20)
        self.header_subtitle = tk.Label(self.header_frame, text="Manage your data processing tools",
                                        font=("Segoe UI", 10), bg=self.bg_header,
                                        fg=self.text_light, padx=10)
        self.header_subtitle.pack(side="left", pady=20)

        self.main_screen = tk.Frame(self.right_container, bg=self.bg_main, padx=30, pady=20)
        self.main_screen.pack(side="top", fill="both", expand=True)

    # --------------------------------------------------
    # Navigation
    # --------------------------------------------------
    def clear_main_screen(self):
        for widget in self.main_screen.winfo_children():
            widget.destroy()
        self.header_title.config(text="Dashboard Overview")
        self.header_subtitle.config(text="Manage your data processing tools")

    def update_header(self, title, subtitle):
        self.header_title.config(text=title)
        self.header_subtitle.config(text=subtitle)

    def show_dashboard_home(self):
        self.clear_main_screen()
        self.main_screen.columnconfigure(0, weight=1)
        self.main_screen.columnconfigure(1, weight=1)
        self.main_screen.rowconfigure(0, weight=1)
        self.main_screen.rowconfigure(1, weight=1)

        self.create_card(self.main_screen, 0, 0, "🧹", "Data Cleaner",
                         "Remove empty or null rows from Excel/CSV files.",
                         "Launch Cleaner", self.load_data_cleaner)
        self.create_card(self.main_screen, 0, 1, "📄", "Attendance PDF Generator",
                         "Generate per-source PDF reports.",
                         "Launch Generator", lambda: self.load_placeholder("PDF Generator"))
        self.create_card(self.main_screen, 1, 0, "📋", "Source Adder",
                         "Enrich attendance lists with sponsorship data.",
                         "Launch Adder", self.load_source_adder)
        self.create_card(self.main_screen, 1, 1, "🚫", "Remove Coordinators",
                         "Strip coordinator rows and empty-value rows from attendance files.",
                         "Launch Remover", self.load_remove_coordinators)

    def load_tool(self, ToolClass, title, subtitle):
        self.clear_main_screen()
        self.update_header(title, subtitle)
        tool_instance = ToolClass(parent=self.main_screen, controller=self)
        tool_instance.pack(fill="both", expand=True)

    def load_deduplicator(self):
        self.load_tool(DeduplicateTool, "Remove Duplicate Name", "Attendance Deduplicator")

    def load_email_phone_merger(self):
        self.load_tool(EmailPhoneMergerTool, "Add Email and Phone", "Attendance + Registration Merger")

    def load_remove_coordinators(self):
        self.load_tool(RemoveCoordinatorsTool, "Remove Coordinators",
                       "Filter out coordinator and empty-value rows")

    def load_data_cleaner(self):
        self.load_tool(DataCleanerTool, "Data Cleaner & Null Remover", "Clean empty or null rows")

    def load_source_adder(self):
        self.load_tool(SourceAdderTool, "Attendance Source Adder",
                       "Match sponsorships via ID, Email or Phone")

    def load_placeholder(self, name):
        self.load_tool(PlaceholderTool, name, "Module coming soon...")

    # --------------------------------------------------
    # UI helpers
    # --------------------------------------------------
    def create_collapsible_folder(self, parent, title, sub_items):
        folder_frame = tk.Frame(parent, bg=self.bg_sidebar)
        folder_frame.pack(fill="x", pady=2, padx=10)

        toggle_btn = ttk.Button(folder_frame, text=f"📁 {title} ▶",
                                style="SidebarFolder.TButton")
        toggle_btn.pack(fill="x")
        content_frame = tk.Frame(folder_frame, bg=self.bg_sidebar)

        def toggle():
            if content_frame.winfo_ismapped():
                content_frame.pack_forget()
                toggle_btn.config(text=f"📁 {title} ▶")
            else:
                content_frame.pack(fill="x", pady=(5, 0))
                toggle_btn.config(text=f"📁 {title} ▼")

        toggle_btn.config(command=toggle)

        for item_name, command in sub_items:
            btn = ttk.Button(content_frame, text=f"• {item_name}",
                             style="SidebarItem.TButton", command=command)
            btn.pack(fill="x", padx=(20, 0), pady=1)

    def create_card(self, parent, row, col, icon, title, description, btn_text, command):
        card = tk.Frame(parent, bg="white", bd=1, relief="solid",
                        highlightbackground=self.border_color, highlightthickness=1)
        card.grid(row=row, column=col, padx=10, pady=10, sticky="nsew")
        inner = tk.Frame(card, bg="white", padx=25, pady=25)
        inner.pack(fill="both", expand=True)
        tk.Label(inner, text=icon, font=("Segoe UI", 32), bg="white", fg=self.text_dark).pack(anchor="w")
        tk.Label(inner, text=title, font=("Segoe UI", 15, "bold"), bg="white",
                 fg=self.text_dark).pack(anchor="w", pady=(10, 5))
        tk.Label(inner, text=description, font=("Segoe UI", 10), bg="white",
                 fg=self.text_light, wraplength=350, justify="left").pack(anchor="w", fill="x", expand=True)
        ttk.Button(inner, text=btn_text, style="Launch.TButton",
                   command=command).pack(anchor="w", pady=(15, 0))

    def create_placeholder_card(self, parent, row, col):
        card = tk.Frame(parent, bg="#f8fafc", bd=1, relief="solid",
                        highlightbackground=self.border_color, highlightthickness=1)
        card.grid(row=row, column=col, padx=10, pady=10, sticky="nsew")
        inner = tk.Frame(card, bg="#f8fafc", padx=25, pady=25)
        inner.pack(fill="both", expand=True)
        tk.Label(inner, text="➕", font=("Segoe UI", 32), bg="#f8fafc",
                 fg="#cbd5e1").pack(anchor="w")
        tk.Label(inner, text="More Tools Coming Soon", font=("Segoe UI", 15, "bold"),
                 bg="#f8fafc", fg="#94a3b8").pack(anchor="w", pady=(10, 5))


# ==========================================
# TOOL 1: DEDUPLICATOR
# ==========================================
class DeduplicateTool(ttk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, style="TFrame")
        self.controller = controller
        self.file_path = tk.StringVar()
        self.data = None
        self.result = None

        main_frame = ttk.Frame(self, padding="15")
        main_frame.pack(fill="both", expand=True)

        desc_text = (
            "📌 This tool deduplicates attendance records by student name.\n"
            "   • Select an Excel file with columns: Name, Join, Left, Duration.\n"
            "   • Click 'Process' – duplicate names are merged, and durations are summed."
        )
        ttk.Label(main_frame, text=desc_text, font=("Segoe UI", 9),
                  foreground="#2a4b6c", justify="left").pack(anchor="w", pady=(0, 10))
        ttk.Separator(main_frame, orient="horizontal").pack(fill="x", pady=(0, 12))

        file_frame = ttk.Frame(main_frame)
        file_frame.pack(fill="x", pady=(0, 10))
        ttk.Label(file_frame, text="Excel File:").pack(side="left", padx=(0, 8))
        ttk.Entry(file_frame, textvariable=self.file_path, width=55).pack(
            side="left", padx=(0, 8), fill="x", expand=True)
        ttk.Button(file_frame, text="📂 Browse", command=self.browse_file).pack(side="left", padx=(0, 6))
        ttk.Button(file_frame, text="⚡ Process", command=self.process_file,
                   style="Success.TButton").pack(side="left", padx=(0, 6))
        ttk.Button(file_frame, text="🗑️ Clear", command=self.clear_all,
                   style="Danger.TButton").pack(side="left", padx=(0, 6))
        ttk.Button(file_frame, text="💾 Save as Excel", command=self.save_result,
                   style="Primary.TButton").pack(side="left")

        self.status = ttk.Label(main_frame, text="Ready", relief="sunken",
                                anchor="w", style="Status.TLabel")
        self.status.pack(fill="x", pady=(0, 10))

        tree_frame = ttk.Frame(main_frame)
        tree_frame.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(tree_frame, columns=("Name", "Join", "Left", "Duration"),
                                 show="headings", height=18)
        self.tree.heading("Name", text="Name")
        self.tree.heading("Join", text="Join Time")
        self.tree.heading("Left", text="Left Time")
        self.tree.heading("Duration", text="Total Duration (HH:MM:SS)")
        self.tree.column("Name", width=280, anchor="w")
        self.tree.column("Join", width=180, anchor="center")
        self.tree.column("Left", width=180, anchor="center")
        self.tree.column("Duration", width=150, anchor="center")
        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        self.tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        bottom_frame = ttk.Frame(main_frame)
        bottom_frame.pack(fill="x", pady=(12, 0))
        self.count_label = ttk.Label(bottom_frame, text="", font=("Segoe UI", 9, "italic"))
        self.count_label.pack(side="right", padx=(0, 15))

    def browse_file(self):
        filename = filedialog.askopenfilename(title="Select Attendance Excel File",
                                              filetypes=[("Excel files", "*.xlsx *.xls")])
        if filename:
            self.file_path.set(filename)
            self.status.config(text=f"📁 Loaded: {os.path.basename(filename)}")

    def process_file(self):
        if not self.file_path.get():
            messagebox.showerror("Error", "Please select an Excel file first.")
            return
        try:
            self.status.config(text="Processing data... Please wait.")
            self.update_idletasks()
            grouped_df, original_count, corrections = process_duplicates(self.file_path.get())
            self.result = grouped_df
            self.populate_tree(grouped_df)
            msg = f"✅ Processed: {original_count} rows → {len(grouped_df)} unique names"
            if corrections > 0:
                msg += f" (Corrected {corrections} durations)"
            self.status.config(text=msg)
            self.count_label.config(text=f"{len(grouped_df)} unique entries")
            messagebox.showinfo("Success",
                                f"Deduplication complete.\nOriginal: {original_count} rows\n"
                                f"Unique: {len(grouped_df)} rows")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to process file:\n{str(e)}")
            self.status.config(text="❌ Error")

    def populate_tree(self, df):
        for item in self.tree.get_children():
            self.tree.delete(item)
        for _, row in df.iterrows():
            self.tree.insert("", "end", values=list(row))

    def save_result(self):
        if self.result is None or self.result.empty:
            messagebox.showwarning("Warning", "No data to save. Process a file first.")
            return
        file_path = filedialog.asksaveasfilename(title="Save Deduplicated Excel",
                                                 defaultextension=".xlsx",
                                                 filetypes=[("Excel files", "*.xlsx")])
        if file_path:
            try:
                wb = Workbook()
                ws = wb.active
                ws.title = "Deduplicated"
                headers = ["Name", "Join", "Left", "Duration"]
                for col_idx, hdr in enumerate(headers, 1):
                    ws.cell(row=1, column=col_idx, value=hdr)
                for row_idx, row in enumerate(self.result.itertuples(index=False), start=2):
                    for col_idx, val in enumerate(row, 1):
                        ws.cell(row=row_idx, column=col_idx, value=val)
                for col_idx in range(1, len(headers) + 1):
                    col_letter = get_column_letter(col_idx)
                    max_len = max(len(headers[col_idx-1]), 10)
                    for row_idx in range(2, len(self.result)+2):
                        val = ws.cell(row=row_idx, column=col_idx).value
                        if val:
                            max_len = max(max_len, len(str(val)))
                    ws.column_dimensions[col_letter].width = min(max_len + 2, 40)
                wb.save(file_path)
                messagebox.showinfo("Success", f"File saved to:\n{file_path}")
                self.status.config(text=f"💾 Saved: {os.path.basename(file_path)}")
            except Exception as e:
                messagebox.showerror("Error", f"Could not save file:\n{str(e)}")

    def clear_all(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.data = None
        self.result = None
        self.file_path.set("")
        self.status.config(text="Cleared")
        self.count_label.config(text="")


# ==========================================
# TOOL 2: EMAIL & PHONE MERGER
# ==========================================
class EmailPhoneMergerTool(ttk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, style="TFrame")
        self.controller = controller
        self.reg_path = tk.StringVar()
        self.att_path = tk.StringVar()
        self.result_data = []
        self.unmatched_names = []

        main_frame = ttk.Frame(self, padding="15")
        main_frame.pack(fill="both", expand=True)

        desc_text = (
            "📌 This tool merges attendance with registration details.\n"
            "   • Select the Registration file (columns: Name, Email, IDNumber, Phone).\n"
            "   • Select the Attendance file (columns: Name, Join, Left, Duration).\n"
            "   • It matches by name (handling slight variations) to add Email & Phone."
        )
        ttk.Label(main_frame, text=desc_text, font=("Segoe UI", 9), foreground="#2a4b6c",
                  justify="left", wraplength=900).pack(anchor="w", pady=(0, 10))
        ttk.Separator(main_frame, orient="horizontal").pack(fill="x", pady=(0, 10))

        reg_frame = ttk.Frame(main_frame)
        reg_frame.pack(fill="x", pady=3)
        ttk.Label(reg_frame, text="📁 Registration File:").pack(side="left", padx=(0, 8))
        ttk.Entry(reg_frame, textvariable=self.reg_path, width=60).pack(
            side="left", fill="x", expand=True, padx=(0, 8))
        ttk.Button(reg_frame, text="Browse",
                   command=lambda: self.browse_file(self.reg_path)).pack(side="left")

        att_frame = ttk.Frame(main_frame)
        att_frame.pack(fill="x", pady=3)
        ttk.Label(att_frame, text="📁 Attendance File:").pack(side="left", padx=(0, 8))
        ttk.Entry(att_frame, textvariable=self.att_path, width=60).pack(
            side="left", fill="x", expand=True, padx=(0, 8))
        ttk.Button(att_frame, text="Browse",
                   command=lambda: self.browse_file(self.att_path)).pack(side="left")

        btn_frame = ttk.Frame(main_frame)
        btn_frame.pack(fill="x", pady=10)
        ttk.Button(btn_frame, text="⚡ Merge & Match", command=self.process_files,
                   style="Success.TButton").pack(side="left", padx=(0, 10))
        ttk.Button(btn_frame, text="Clear All", command=self.clear_all,
                   style="Danger.TButton").pack(side="left", padx=(0, 10))
        ttk.Button(btn_frame, text="💾 Save as Excel", command=self.save_result,
                   style="Primary.TButton").pack(side="left")

        self.status = ttk.Label(main_frame, text="Ready", relief="sunken",
                                anchor="w", style="Status.TLabel")
        self.status.pack(fill="x", pady=(5, 10))

        tree_frame = ttk.Frame(main_frame)
        tree_frame.pack(fill="both", expand=True)
        columns = ("Name", "Email", "Phone", "IDNumber", "Join", "Left", "Duration")
        self.tree = ttk.Treeview(tree_frame, columns=columns, show="headings", height=18)
        col_widths = {"Name": 200, "Email": 200, "Phone": 120, "IDNumber": 150,
                      "Join": 150, "Left": 150, "Duration": 100}
        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=col_widths.get(col, 120),
                             anchor="w" if col == "Name" else "center")
        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        self.tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        bottom = ttk.Frame(main_frame)
        bottom.pack(fill="x", pady=10)
        self.count_label = ttk.Label(bottom, text="", font=("Segoe UI", 9, "italic"))
        self.count_label.pack(side="right", padx=(0, 15))

    def browse_file(self, var):
        f = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xlsm"),
                                                  ("All files", "*.*")])
        if f:
            var.set(f)

    def process_files(self):
        if not self.reg_path.get() or not self.att_path.get():
            messagebox.showerror("Error", "Please select both files.")
            return
        try:
            self.status.config(text="Processing... Please wait.")
            self.update_idletasks()
            self.result_data, self.unmatched_names, matched_count, total = merge_data(
                self.reg_path.get(), self.att_path.get())
            self.populate_tree()
            self.status.config(
                text=f"✅ Matching complete: {matched_count}/{total} matched. "
                     f"Unmatched: {len(self.unmatched_names)}")
            self.count_label.config(text=f"{total} records, {matched_count} matched")
            if self.unmatched_names:
                msg = "The following attendance names could NOT be matched:\n" + \
                      "\n".join(self.unmatched_names[:20])
                if len(self.unmatched_names) > 20:
                    msg += f"\n... and {len(self.unmatched_names)-20} more."
                messagebox.showwarning("Unmatched Names", msg)
        except Exception as e:
            messagebox.showerror("Error", f"Processing failed:\n{str(e)}")
            self.status.config(text="❌ Error")

    def populate_tree(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        for row in self.result_data:
            values = (row["Name"], row["Email"], row["Phone"], row["IDNumber"],
                      row["Join"], row["Left"], row["Duration"])
            self.tree.insert("", "end", values=values)

    def save_result(self):
        if not self.result_data:
            messagebox.showwarning("Warning", "No data to save. Process files first.")
            return
        file_path = filedialog.asksaveasfilename(title="Save Merged Excel",
                                                 defaultextension=".xlsx",
                                                 filetypes=[("Excel files", "*.xlsx")])
        if file_path:
            try:
                wb = Workbook()
                ws = wb.active
                ws.title = "Merged"
                headers = ["Name", "Email", "Phone", "IDNumber", "Join", "Left", "Duration"]
                for col_idx, hdr in enumerate(headers, 1):
                    ws.cell(row=1, column=col_idx, value=hdr)
                for row_idx, row in enumerate(self.result_data, start=2):
                    ws.cell(row=row_idx, column=1, value=row["Name"])
                    ws.cell(row=row_idx, column=2, value=row["Email"])
                    ws.cell(row=row_idx, column=3, value=row["Phone"])
                    ws.cell(row=row_idx, column=4, value=row["IDNumber"])
                    ws.cell(row=row_idx, column=5, value=row["Join"])
                    ws.cell(row=row_idx, column=6, value=row["Left"])
                    ws.cell(row=row_idx, column=7, value=row["Duration"])
                for col_idx in range(1, len(headers) + 1):
                    col_letter = get_column_letter(col_idx)
                    max_len = max(len(headers[col_idx-1]), 10)
                    for row_idx in range(2, len(self.result_data)+2):
                        val = ws.cell(row=row_idx, column=col_idx).value
                        if val:
                            max_len = max(max_len, len(str(val)))
                    ws.column_dimensions[col_letter].width = min(max_len + 2, 40)
                wb.save(file_path)
                messagebox.showinfo("Success", f"File saved to:\n{file_path}")
                self.status.config(text=f"💾 Saved: {os.path.basename(file_path)}")
            except Exception as e:
                messagebox.showerror("Error", f"Could not save file:\n{str(e)}")

    def clear_all(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.result_data = []
        self.unmatched_names = []
        self.reg_path.set("")
        self.att_path.set("")
        self.status.config(text="Cleared")
        self.count_label.config(text="")


# ==========================================
# TOOL 3: DATA CLEANER
# ==========================================
class DataCleanerTool(ttk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, style="TFrame")
        self.controller = controller
        self.file_path = tk.StringVar()
        self.selected_column = tk.StringVar()
        self.headers = []
        self.kept_rows = []
        self.removed_rows = []

        main_frame = ttk.Frame(self, padding="15")
        main_frame.pack(fill="both", expand=True)

        desc_text = (
            "📌 This tool removes rows where a specific column contains empty, null, or whitespace-only values.\n"
            "   • Select an Excel or CSV file and choose the column to check.\n"
            "   • The grid below will show ONLY the removed rows.\n"
            "   • Click 'Save Cleaned Result' to save the remaining rows as a new Excel file."
        )
        ttk.Label(main_frame, text=desc_text, font=("Segoe UI", 9),
                  foreground="#2a4b6c", justify="left").pack(anchor="w", pady=(0, 10))
        ttk.Separator(main_frame, orient="horizontal").pack(fill="x", pady=(0, 12))

        file_frame = ttk.Frame(main_frame)
        file_frame.pack(fill="x", pady=(0, 10))
        ttk.Label(file_frame, text="Data File:").pack(side="left", padx=(0, 8))
        ttk.Entry(file_frame, textvariable=self.file_path, width=35).pack(
            side="left", padx=(0, 8), fill="x", expand=True)
        ttk.Button(file_frame, text="📂 Browse", command=self.browse_file).pack(side="left", padx=(0, 6))
        ttk.Label(file_frame, text="Check Column:").pack(side="left", padx=(10, 5))
        self.col_combo = ttk.Combobox(file_frame, textvariable=self.selected_column,
                                      state="disabled", width=22)
        self.col_combo.pack(side="left", padx=(0, 10))
        ttk.Button(file_frame, text="⚡ Process", command=self.process_file,
                   style="Success.TButton").pack(side="left", padx=(0, 6))
        ttk.Button(file_frame, text="🗑️ Clear", command=self.clear_all,
                   style="Danger.TButton").pack(side="left", padx=(0, 6))
        ttk.Button(file_frame, text="💾 Save Cleaned Result", command=self.save_result,
                   style="Primary.TButton").pack(side="left")

        self.status = ttk.Label(main_frame, text="Ready. Please select a file.",
                                relief="sunken", anchor="w", style="Status.TLabel")
        self.status.pack(fill="x", pady=(10, 10))

        self.tree_title_label = ttk.Label(main_frame, text="⚠️ Removed Rows Preview:",
                                          font=("Segoe UI", 10, "bold"), foreground="#dc3545")
        self.tree_title_label.pack(anchor="w", pady=(0, 5))

        tree_frame = ttk.Frame(main_frame)
        tree_frame.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(tree_frame, show="headings", height=16)
        yscrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        xscrollbar = ttk.Scrollbar(tree_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=yscrollbar.set, xscrollcommand=xscrollbar.set)
        yscrollbar.pack(side="right", fill="y")
        xscrollbar.pack(side="bottom", fill="x")
        self.tree.pack(side="left", fill="both", expand=True)

        bottom_frame = ttk.Frame(main_frame)
        bottom_frame.pack(fill="x", pady=(12, 0))
        self.count_label = ttk.Label(bottom_frame, text="", font=("Segoe UI", 9, "italic"))
        self.count_label.pack(side="right", padx=(0, 15))

    def browse_file(self):
        filename = filedialog.askopenfilename(
            title="Select Data File",
            filetypes=[("Excel & CSV", "*.csv *.xlsx *.xlsm"),
                       ("Excel files", "*.xlsx *.xlsm"), ("CSV files", "*.csv")])
        if filename:
            self.file_path.set(filename)
            self.status.config(text=f"📁 Loaded: {os.path.basename(filename)}")
            self.load_columns(filename)

    def load_columns(self, filename):
        try:
            headers = get_headers_dc(filename)
            self.col_combo['values'] = ["ALL COLUMNS (Check entire row)"] + headers
            self.col_combo.current(0)
            self.col_combo.config(state="readonly")
            self.tree.delete(*self.tree.get_children())
            self.tree["columns"] = []
            self.count_label.config(text="")
        except Exception as e:
            messagebox.showerror("Error", f"Could not read file headers:\n{str(e)}")
            self.col_combo.config(state="disabled")

    def process_file(self):
        if not self.file_path.get() or not self.selected_column.get():
            messagebox.showerror("Error", "Please select a file and a column first.")
            return
        try:
            self.status.config(text="Processing... Please wait.")
            self.update_idletasks()
            self.headers, self.kept_rows, self.removed_rows = process_file_data(
                self.file_path.get(), self.selected_column.get())
            self.populate_tree()
            self.status.config(text=f"✅ Done: Removed {len(self.removed_rows)} rows. "
                                    f"Remaining: {len(self.kept_rows)}")
            self.count_label.config(text=f"{len(self.removed_rows)} rows removed")
            messagebox.showinfo("Summary",
                                f"Process Complete!\n\n"
                                f"Rows Removed: {len(self.removed_rows)}\n"
                                f"Remaining Rows: {len(self.kept_rows)}\n\n"
                                f"Click 'Save Cleaned Result' to save the remaining rows.")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to process file:\n{str(e)}")
            self.status.config(text="❌ Error")

    def populate_tree(self):
        self.tree.delete(*self.tree.get_children())
        if not self.removed_rows:
            self.tree["columns"] = ["Status"]
            self.tree.heading("Status", text="No empty rows found!")
            self.tree.column("Status", width=300, anchor="center")
            return
        cols = self.headers
        self.tree["columns"] = cols
        for col in cols:
            self.tree.heading(col, text=col)
            width = 250 if self.selected_column.get() == col else 150
            self.tree.column(col, width=width, anchor="w", stretch=True)
        for row in self.removed_rows[:1000]:
            clean_row = [str(val) if val is not None else "" for val in row]
            while len(clean_row) < len(cols):
                clean_row.append("")
            self.tree.insert("", "end", values=clean_row)

    def save_result(self):
        if not self.kept_rows:
            messagebox.showwarning("Warning", "No cleaned data to save. Process a file first.")
            return
        file_path = filedialog.asksaveasfilename(
            title="Save Cleaned Data", defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")])
        if file_path:
            try:
                wb = Workbook()
                ws = wb.active
                ws.title = "Enriched Attendance"
                for col_idx, hdr in enumerate(self.headers, 1):
                    ws.cell(row=1, column=col_idx, value=hdr)
                for row_idx, row in enumerate(self.kept_rows, start=2):
                    for col_idx, val in enumerate(row, 1):
                        ws.cell(row=row_idx, column=col_idx, value=val)
                for col_idx in range(1, len(self.headers) + 1):
                    col_letter = get_column_letter(col_idx)
                    max_len = max(len(str(self.headers[col_idx-1])), 10)
                    for row_idx in range(2, len(self.kept_rows)+2):
                        val = ws.cell(row=row_idx, column=col_idx).value
                        if val:
                            max_len = max(max_len, len(str(val)))
                    ws.column_dimensions[col_letter].width = min(max_len + 2, 40)
                wb.save(file_path)
                messagebox.showinfo("Success", f"File saved successfully to:\n{file_path}")
                self.status.config(text=f"💾 Saved: {os.path.basename(file_path)}")
            except Exception as e:
                messagebox.showerror("Error", f"Could not save file:\n{str(e)}")

    def clear_all(self):
        self.tree.delete(*self.tree.get_children())
        self.tree["columns"] = []
        self.headers = []
        self.kept_rows = []
        self.removed_rows = []
        self.file_path.set("")
        self.selected_column.set("")
        self.col_combo['values'] = []
        self.col_combo.config(state="disabled")
        self.status.config(text="Cleared")
        self.count_label.config(text="")


# ==========================================
# TOOL 4: SOURCE ADDER
# ==========================================
class SourceAdderTool(ttk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, style="TFrame")
        self.controller = controller
        self.sponsor_path = tk.StringVar()
        self.att_path = tk.StringVar()
        self.result_data = []
        self.unmatched_count = 0

        main_frame = ttk.Frame(self, padding="15")
        main_frame.pack(fill="both", expand=True)

        desc_text = (
            "📌 This tool adds a 'Source' column to your attendance list.\n"
            "   • Matches via IDNumber, with a fallback to Email or Phone.\n"
            "   • Unmatched records are reported so you can manually correct them."
        )
        ttk.Label(main_frame, text=desc_text, font=("Segoe UI", 9), foreground="#2a4b6c",
                  justify="left", wraplength=1000).pack(anchor="w", pady=(0, 10))
        ttk.Separator(main_frame, orient="horizontal").pack(fill="x", pady=(0, 10))

        sponsor_frame = ttk.Frame(main_frame)
        sponsor_frame.pack(fill="x", pady=3)
        ttk.Label(sponsor_frame, text="📁 Sponsorship File:").pack(side="left", padx=(0, 8))
        ttk.Entry(sponsor_frame, textvariable=self.sponsor_path, width=60).pack(
            side="left", fill="x", expand=True, padx=(0, 8))
        ttk.Button(sponsor_frame, text="Browse",
                   command=lambda: self.browse_file(self.sponsor_path)).pack(side="left")

        att_frame = ttk.Frame(main_frame)
        att_frame.pack(fill="x", pady=3)
        ttk.Label(att_frame, text="📁 Attendance File:").pack(side="left", padx=(0, 8))
        ttk.Entry(att_frame, textvariable=self.att_path, width=60).pack(
            side="left", fill="x", expand=True, padx=(0, 8))
        ttk.Button(att_frame, text="Browse",
                   command=lambda: self.browse_file(self.att_path)).pack(side="left")

        btn_frame = ttk.Frame(main_frame)
        btn_frame.pack(fill="x", pady=10)
        ttk.Button(btn_frame, text="⚡ Add Sources", command=self.process_files,
                   style="Success.TButton").pack(side="left", padx=(0, 10))
        ttk.Button(btn_frame, text="Clear All", command=self.clear_all,
                   style="Danger.TButton").pack(side="left", padx=(0, 10))
        ttk.Button(btn_frame, text="💾 Save as Excel", command=self.save_result,
                   style="Primary.TButton").pack(side="left")

        self.status = ttk.Label(main_frame, text="Ready", relief="sunken",
                                anchor="w", style="Status.TLabel")
        self.status.pack(fill="x", pady=(5, 10))

        tree_frame = ttk.Frame(main_frame)
        tree_frame.pack(fill="both", expand=True)
        columns = ("Name", "Email", "Phone", "Source", "IDNumber", "Join", "Left", "Duration")
        self.tree = ttk.Treeview(tree_frame, columns=columns, show="headings", height=18)
        col_widths = {"Name": 180, "Email": 200, "Phone": 120, "Source": 200,
                      "IDNumber": 150, "Join": 130, "Left": 130, "Duration": 100}
        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=col_widths.get(col, 120),
                             anchor="w" if col == "Name" else "center")
        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        self.tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        bottom = ttk.Frame(main_frame)
        bottom.pack(fill="x", pady=10)
        self.count_label = ttk.Label(bottom, text="", font=("Segoe UI", 9, "italic"))
        self.count_label.pack(side="right", padx=(0, 15))

    def browse_file(self, var):
        f = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xlsm"),
                                                  ("All files", "*.*")])
        if f:
            var.set(f)

    def process_files(self):
        if not self.sponsor_path.get() or not self.att_path.get():
            messagebox.showerror("Error", "Please select both files.")
            return
        try:
            self.status.config(text="Processing... Please wait.")
            self.update_idletasks()
            self.result_data, unmatched_names, matched_count, total = match_sources(
                self.sponsor_path.get(), self.att_path.get())
            self.populate_tree()
            self.status.config(
                text=f"✅ Matching complete: {matched_count}/{total} matched. "
                     f"Unmatched: {len(unmatched_names)}")
            self.count_label.config(text=f"{total} attendance records, {matched_count} have source")
            if unmatched_names:
                msg = "The following attendance records could NOT be matched:\n" + \
                      "\n".join(unmatched_names[:20])
                if len(unmatched_names) > 20:
                    msg += f"\n... and {len(unmatched_names)-20} more."
                messagebox.showwarning("Unmatched Names", msg)
        except Exception as e:
            messagebox.showerror("Error", f"Processing failed:\n{str(e)}")
            self.status.config(text="❌ Error")

    def populate_tree(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        for row in self.result_data:
            values = (row["Name"], row["Email"], row["Phone"], row["Source"],
                      row["IDNumber"], row["Join"], row["Left"], row["Duration"])
            self.tree.insert("", "end", values=values)

    def save_result(self):
        if not self.result_data:
            messagebox.showwarning("Warning", "No data to save. Process files first.")
            return
        file_path = filedialog.asksaveasfilename(
            title="Save Enriched Attendance", defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")])
        if file_path:
            try:
                wb = Workbook()
                ws = wb.active
                ws.title = "Enriched Attendance"
                headers = ["Name", "Email", "Phone", "Source", "IDNumber", "Join", "Left", "Duration"]
                for col_idx, hdr in enumerate(headers, 1):
                    ws.cell(row=1, column=col_idx, value=hdr)
                for row_idx, row in enumerate(self.result_data, start=2):
                    ws.cell(row=row_idx, column=1, value=row["Name"])
                    ws.cell(row=row_idx, column=2, value=row["Email"])
                    ws.cell(row=row_idx, column=3, value=row["Phone"])
                    ws.cell(row=row_idx, column=4, value=row["Source"])
                    ws.cell(row=row_idx, column=5, value=row["IDNumber"])
                    ws.cell(row=row_idx, column=6, value=row["Join"])
                    ws.cell(row=row_idx, column=7, value=row["Left"])
                    ws.cell(row=row_idx, column=8, value=row["Duration"])
                for col_idx in range(1, 9):
                    col_letter = get_column_letter(col_idx)
                    max_len = max(len(headers[col_idx-1]), 10)
                    for row_idx in range(2, len(self.result_data)+2):
                        val = ws.cell(row=row_idx, column=col_idx).value
                        if val:
                            max_len = max(max_len, len(str(val)))
                    ws.column_dimensions[col_letter].width = min(max_len + 2, 40)
                wb.save(file_path)
                messagebox.showinfo("Success", f"File saved to:\n{file_path}")
                self.status.config(text=f"💾 Saved: {os.path.basename(file_path)}")
            except Exception as e:
                messagebox.showerror("Error", f"Could not save file:\n{str(e)}")

    def clear_all(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.result_data = []
        self.sponsor_path.set("")
        self.att_path.set("")
        self.status.config(text="Cleared")
        self.count_label.config(text="")


# ==========================================
# TOOL 5: REMOVE COORDINATORS + EMPTY ROWS
# ==========================================
class RemoveCoordinatorsTool(ttk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, style="TFrame")
        self.controller = controller

        # ---- state vars ----
        self.file_path        = tk.StringVar()
        self.match_mode       = tk.StringVar(value="role_column")
        self.role_column      = tk.StringVar()
        self.keywords_text    = tk.StringVar(value="coordinator, coord, lead, admin")
        self.remove_empty_var = tk.BooleanVar(value=True)
        self.empty_column     = tk.StringVar(value="ALL COLUMNS (Check entire row)")

        self.headers         = []
        self.kept_rows       = []
        self.removed_rows    = []
        self.removed_reasons = []

        # ---- layout ----
        main_frame = ttk.Frame(self, padding="15")
        main_frame.pack(fill="both", expand=True)

        desc_text = (
            "📌 Removes coordinator rows AND/OR rows with empty values.\n"
            "   • Empty-row rule: drop the whole row if a chosen column is blank\n"
            "     (or use 'ALL COLUMNS' to drop a row if ANY cell is blank).\n"
            "   • Coordinator rule: match a role/designation column or a manual list.\n"
            "   • Preview lists only the REMOVED rows, with a reason per row."
        )
        ttk.Label(main_frame, text=desc_text, font=("Segoe UI", 9),
                  foreground="#2a4b6c", justify="left",
                  wraplength=1000).pack(anchor="w", pady=(0, 10))
        ttk.Separator(main_frame, orient="horizontal").pack(fill="x", pady=(0, 10))

        # --- File row ---
        file_frame = ttk.Frame(main_frame)
        file_frame.pack(fill="x", pady=(0, 10))
        ttk.Label(file_frame, text="Data File:").pack(side="left", padx=(0, 8))
        ttk.Entry(file_frame, textvariable=self.file_path, width=40).pack(
            side="left", padx=(0, 8), fill="x", expand=True)
        ttk.Button(file_frame, text="📂 Browse",
                   command=self.browse_file).pack(side="left", padx=(0, 6))
        ttk.Button(file_frame, text="⚡ Process",
                   command=self.process_file,
                   style="Success.TButton").pack(side="left", padx=(0, 6))
        ttk.Button(file_frame, text="🗑️ Clear",
                   command=self.clear_all,
                   style="Danger.TButton").pack(side="left", padx=(0, 6))
        ttk.Button(file_frame, text="💾 Save Cleaned Result",
                   command=self.save_result,
                   style="Primary.TButton").pack(side="left")

        # ---- Rules frame ----
        rules = ttk.LabelFrame(main_frame, text="Removal Rules", padding=10)
        rules.pack(fill="x", pady=(0, 10))

        # Empty-row rule
        empty_row = ttk.Frame(rules)
        empty_row.pack(fill="x", pady=(0, 8))
        ttk.Checkbutton(
            empty_row,
            text="Remove rows where a column is empty / null",
            variable=self.remove_empty_var,
            command=self._on_empty_toggle,
        ).pack(side="left")
        ttk.Label(empty_row, text="  Check Column:").pack(side="left", padx=(15, 6))
        self.empty_combo = ttk.Combobox(
            empty_row, textvariable=self.empty_column,
            state="disabled", width=32)
        self.empty_combo.pack(side="left")

        ttk.Separator(rules, orient="horizontal").pack(fill="x", pady=8)

        # Coordinator mode radio buttons
        mode_row = ttk.Frame(rules)
        mode_row.pack(fill="x", pady=(0, 8))
        ttk.Label(mode_row, text="Coordinator mode:",
                  font=("Segoe UI", 10, "bold")).pack(side="left", padx=(0, 10))
        ttk.Radiobutton(mode_row, text="Role Column", value="role_column",
                        variable=self.match_mode,
                        command=self._on_mode_change).pack(side="left", padx=(0, 10))
        ttk.Radiobutton(mode_row, text="Manual List", value="manual_list",
                        variable=self.match_mode,
                        command=self._on_mode_change).pack(side="left", padx=(0, 10))
        ttk.Radiobutton(mode_row, text="Both", value="both",
                        variable=self.match_mode,
                        command=self._on_mode_change).pack(side="left")

        role_row = ttk.Frame(rules)
        role_row.pack(fill="x", pady=(0, 6))
        ttk.Label(role_row, text="Role Column:").pack(side="left", padx=(0, 6))
        self.role_combo = ttk.Combobox(role_row, textvariable=self.role_column,
                                       state="disabled", width=25)
        self.role_combo.pack(side="left", padx=(0, 15))
        ttk.Label(role_row, text="Keywords (comma-separated):").pack(side="left", padx=(0, 6))
        ttk.Entry(role_row, textvariable=self.keywords_text, width=45).pack(
            side="left", fill="x", expand=True)

        manual_row = ttk.Frame(rules)
        manual_row.pack(fill="x", pady=(0, 4))
        ttk.Label(manual_row,
                  text="Manual Coordinators (one per line — Name / Email / Phone / ID):"
                  ).pack(anchor="w")
        self.manual_box = tk.Text(manual_row, height=4, font=("Segoe UI", 9),
                                  wrap="word", bd=1, relief="solid")
        self.manual_box.pack(fill="x", pady=(4, 0))

        # ---- Status ----
        self.status = ttk.Label(main_frame, text="Ready. Please select a file.",
                                relief="sunken", anchor="w", style="Status.TLabel")
        self.status.pack(fill="x", pady=(10, 10))

        # ---- Preview ----
        self.tree_title_label = ttk.Label(
            main_frame, text="⚠️ Removed Rows Preview:",
            font=("Segoe UI", 10, "bold"), foreground="#dc3545")
        self.tree_title_label.pack(anchor="w", pady=(0, 5))

        tree_frame = ttk.Frame(main_frame)
        tree_frame.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(tree_frame, show="headings", height=14)
        yscroll = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        xscroll = ttk.Scrollbar(tree_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        yscroll.pack(side="right", fill="y")
        xscroll.pack(side="bottom", fill="x")
        self.tree.pack(side="left", fill="both", expand=True)

        bottom_frame = ttk.Frame(main_frame)
        bottom_frame.pack(fill="x", pady=(12, 0))
        self.count_label = ttk.Label(bottom_frame, text="", font=("Segoe UI", 9, "italic"))
        self.count_label.pack(side="right", padx=(0, 15))

        # Initial enable/disable
        self._on_mode_change()
        self._on_empty_toggle()

    # ---------- enable/disable helpers ----------
    def _on_mode_change(self):
        mode = self.match_mode.get()
        role_on = mode in ("role_column", "both")
        manual_on = mode in ("manual_list", "both")
        self.role_combo.config(state="readonly" if role_on else "disabled")
        self.manual_box.config(state="normal" if manual_on else "disabled",
                               bg="white" if manual_on else "#f0f0f0")

    def _on_empty_toggle(self):
        on = self.remove_empty_var.get()
        self.empty_combo.config(state="readonly" if on else "disabled")

    # ---------- file IO ----------
    def browse_file(self):
        filename = filedialog.askopenfilename(
            title="Select Data File",
            filetypes=[("Excel & CSV", "*.csv *.xlsx *.xlsm"),
                       ("Excel files", "*.xlsx *.xlsm"),
                       ("CSV files", "*.csv")])
        if filename:
            self.file_path.set(filename)
            self.status.config(text=f"📁 Loaded: {os.path.basename(filename)}")
            self.load_columns(filename)

    def load_columns(self, filename):
        try:
            headers = get_file_headers(filename)
            self.headers = headers

            # Role column combo
            self.role_combo['values'] = headers
            self.role_combo.config(state="readonly")
            if headers:
                for h in headers:
                    if h.lower() in ("role", "designation", "user type", "user_type", "type"):
                        self.role_column.set(h)
                        break
                else:
                    self.role_column.set(headers[0])

            # Empty column combo
            self.empty_combo['values'] = ["ALL COLUMNS (Check entire row)"] + headers
            self.empty_combo.current(0)
            self.empty_combo.config(state="readonly")

            self.tree.delete(*self.tree.get_children())
            self.tree["columns"] = []
            self.count_label.config(text="")
        except Exception as e:
            messagebox.showerror("Error", f"Could not read file headers:\n{str(e)}")
            self.role_combo.config(state="disabled")
            self.empty_combo.config(state="disabled")

    # ---------- main processing ----------
    def process_file(self):
        if not self.file_path.get():
            messagebox.showerror("Error", "Please select a file first.")
            return

        mode = self.match_mode.get()
        remove_empty = self.remove_empty_var.get()
        keywords = [k for k in self.keywords_text.get().split(",") if k.strip()]
        manual_lines = [ln.strip() for ln in self.manual_box.get("1.0", "end").splitlines()
                        if ln.strip()]

        # Validation
        if not remove_empty and mode not in ("role_column", "manual_list", "both"):
            messagebox.showerror("Error", "Enable at least one rule.")
            return
        if mode in ("role_column", "both") and not self.role_column.get():
            messagebox.showerror("Error", "Please choose a Role Column.")
            return
        if mode in ("role_column", "both") and not keywords:
            messagebox.showerror("Error", "Please enter at least one keyword.")
            return
        if mode in ("manual_list", "both") and not manual_lines:
            messagebox.showerror("Error", "Please paste at least one coordinator.")
            return

        try:
            self.status.config(text="Processing... Please wait.")
            self.update_idletasks()

            (self.headers, self.kept_rows, self.removed_rows,
             self.removed_reasons, stats) = remove_coordinators(
                self.file_path.get(),
                mode=mode,
                column=self.role_column.get() if mode in ("role_column", "both") else None,
                keywords=keywords,
                manual_list=manual_lines,
                remove_empty=remove_empty,
                empty_column=self.empty_column.get() if remove_empty else None,
            )

            self.populate_tree()

            msg = (f"✅ Done: {len(self.removed_rows)} removed "
                   f"({stats['empty']} empty, {stats['coordinator']} coordinator). "
                   f"Remaining: {len(self.kept_rows)}")
            self.status.config(text=msg)
            self.count_label.config(text=f"{len(self.removed_rows)} rows removed")

            messagebox.showinfo(
                "Summary",
                f"Process Complete!\n\n"
                f"Empty-value rows removed : {stats['empty']}\n"
                f"Coordinator rows removed : {stats['coordinator']}\n"
                f"Total removed            : {len(self.removed_rows)}\n"
                f"Remaining rows           : {len(self.kept_rows)}\n\n"
                f"Click 'Save Cleaned Result' to save the remaining rows.")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to process file:\n{str(e)}")
            self.status.config(text="❌ Error")

    # ---------- preview ----------
    def populate_tree(self):
        self.tree.delete(*self.tree.get_children())
        if not self.removed_rows:
            self.tree["columns"] = ["Status"]
            self.tree.heading("Status", text="No rows matched the removal rules.")
            self.tree.column("Status", width=300, anchor="center")
            return

        cols = list(self.headers) + ["Removal Reason"]
        self.tree["columns"] = cols
        for col in self.headers:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=150, anchor="w", stretch=True)
        self.tree.heading("Removal Reason", text="Removal Reason")
        self.tree.column("Removal Reason", width=280, anchor="w", stretch=True)

        for row, reason in zip(self.removed_rows[:1000], self.removed_reasons[:1000]):
            clean = [str(v) if v is not None else "" for v in row]
            while len(clean) < len(self.headers):
                clean.append("")
            clean.append(reason)
            self.tree.insert("", "end", values=clean)

    # ---------- save ----------
    def save_result(self):
        if not self.kept_rows:
            messagebox.showwarning("Warning", "No cleaned data to save. Process a file first.")
            return
        file_path = filedialog.asksaveasfilename(
            title="Save Cleaned Data", defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")])
        if not file_path:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Enriched Attendance"

            for c_idx, hdr in enumerate(self.headers, 1):
                ws.cell(row=1, column=c_idx, value=hdr)
            for r_idx, row in enumerate(self.kept_rows, start=2):
                for c_idx, val in enumerate(row, 1):
                    ws.cell(row=r_idx, column=c_idx, value=val)

            for c_idx in range(1, len(self.headers) + 1):
                letter = get_column_letter(c_idx)
                max_len = max(len(str(self.headers[c_idx - 1])), 10)
                for r_idx in range(2, len(self.kept_rows) + 2):
                    val = ws.cell(row=r_idx, column=c_idx).value
                    if val:
                        max_len = max(max_len, len(str(val)))
                ws.column_dimensions[letter].width = min(max_len + 2, 40)

            wb.save(file_path)
            messagebox.showinfo("Success", f"File saved to:\n{file_path}")
            self.status.config(text=f"💾 Saved: {os.path.basename(file_path)}")
        except Exception as e:
            messagebox.showerror("Error", f"Could not save file:\n{str(e)}")

    # ---------- reset ----------
    def clear_all(self):
        self.tree.delete(*self.tree.get_children())
        self.tree["columns"] = []
        self.headers = []
        self.kept_rows = []
        self.removed_rows = []
        self.removed_reasons = []
        self.file_path.set("")
        self.role_column.set("")
        self.role_combo['values'] = []
        self.empty_column.set("ALL COLUMNS (Check entire row)")
        self.empty_combo['values'] = []
        self.manual_box.delete("1.0", "end")
        self.status.config(text="Cleared")
        self.count_label.config(text="")


# ==========================================
# Placeholder
# ==========================================
class PlaceholderTool(ttk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, style="TFrame")
        ttk.Label(self, text="Module Coming Soon",
                  font=("Segoe UI", 16, "bold")).pack(pady=50)
        ttk.Button(self, text="⬅ Back to Dashboard",
                   command=lambda: controller.show_dashboard_home()).pack()


if __name__ == "__main__":
    root = tk.Tk()
    app = MasterDashboardApp(root)
    root.mainloop()