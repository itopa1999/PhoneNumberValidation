import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
import re
import os

# ---------- Helper functions ----------
def clean_phone(phone):
    """Remove spaces, dashes, parentheses; keep digits and commas."""
    if not phone:
        return ""
    return re.sub(r'[^\d,]', '', str(phone).strip())

def parse_phone_list(phone_str):
    """Split phone string by comma, clean each."""
    if not phone_str:
        return []
    cleaned = clean_phone(phone_str)
    parts = [p.strip() for p in cleaned.split(',') if p.strip()]
    return parts

# ---------- Main Application ----------
class SourceAdderApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Attendance Source Adder (by ID)")
        self.root.geometry("1300x700")
        self.root.configure(bg="#f0f4f8")

        self.setup_styles()

        self.sponsor_path = tk.StringVar()
        self.att_path = tk.StringVar()
        self.result_data = []
        self.unmatched_count = 0

        # ----- Main container -----
        main = ttk.Frame(root, padding="15")
        main.pack(fill="both", expand=True)

        # Title
        ttk.Label(main, text="📋 Attendance Source Adder (by ID)",
                  font=("Segoe UI", 16, "bold"), foreground="#1a2b3c").pack(anchor="w")
        ttk.Label(main, text="Add Source from registration using IDNumber (with fallback to Email/Phone for missing IDs)",
                  font=("Segoe UI", 10), foreground="#5a6b7c").pack(anchor="w", pady=(0, 5))

        # Description
        desc_text = (
            "📌 This app adds a 'Source' column to your attendance list.\n"
            "   • If attendance already has an IDNumber, it is used directly to get the Source.\n"
            "   • If IDNumber is missing, the app tries to match by Email or Phone to get both IDNumber and Source.\n"
            "   • The output includes: Name, Email, Phone, Source, IDNumber, Join, Left, Duration.\n"
            "   • Unmatched records are reported – you can then manually correct them."
        )
        desc_label = ttk.Label(main, text=desc_text, font=("Segoe UI", 9),
                               foreground="#2a4b6c", background="#f0f4f8",
                               justify="left", wraplength=1000)
        desc_label.pack(anchor="w", pady=(0, 10))

        ttk.Separator(main, orient="horizontal").pack(fill="x", pady=(0, 10))

        # ----- File selection -----
        sponsor_frame = ttk.Frame(main)
        sponsor_frame.pack(fill="x", pady=3)
        ttk.Label(sponsor_frame, text="📁 Sponsorship File:").pack(side="left", padx=(0, 8))
        ttk.Entry(sponsor_frame, textvariable=self.sponsor_path, width=60).pack(side="left", fill="x", expand=True, padx=(0, 8))
        ttk.Button(sponsor_frame, text="Browse", command=lambda: self.browse_file(self.sponsor_path)).pack(side="left")

        att_frame = ttk.Frame(main)
        att_frame.pack(fill="x", pady=3)
        ttk.Label(att_frame, text="📁 Attendance File:").pack(side="left", padx=(0, 8))
        ttk.Entry(att_frame, textvariable=self.att_path, width=60).pack(side="left", fill="x", expand=True, padx=(0, 8))
        ttk.Button(att_frame, text="Browse", command=lambda: self.browse_file(self.att_path)).pack(side="left")

        # Toolbar with Save button next to Clear
        btn_frame = ttk.Frame(main)
        btn_frame.pack(fill="x", pady=10)

        ttk.Button(btn_frame, text="⚡ Add Sources", command=self.process_files,
                   style="Success.TButton").pack(side="left", padx=(0, 10))
        ttk.Button(btn_frame, text="Clear All", command=self.clear_all,
                   style="Danger.TButton").pack(side="left", padx=(0, 10))
        ttk.Button(btn_frame, text="💾 Save as Excel", command=self.save_result,
                   style="Primary.TButton").pack(side="left")

        # Status bar
        self.status = ttk.Label(main, text="Ready", relief="sunken", anchor="w", style="Status.TLabel")
        self.status.pack(fill="x", pady=(5, 10))

        # Treeview – Source, then IDNumber, then times
        tree_frame = ttk.Frame(main)
        tree_frame.pack(fill="both", expand=True)

        columns = ("Name", "Email", "Phone", "Source", "IDNumber", "Join", "Left", "Duration")
        self.tree = ttk.Treeview(tree_frame, columns=columns, show="headings", height=18)

        col_widths = {"Name": 180, "Email": 200, "Phone": 120, "Source": 200,
                      "IDNumber": 150, "Join": 130, "Left": 130, "Duration": 100}
        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=col_widths.get(col, 120),
                             anchor="w" if col=="Name" else "center")

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        self.tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # Bottom bar (Quit + count label)
        bottom = ttk.Frame(main)
        bottom.pack(fill="x", pady=10)
        ttk.Button(bottom, text="Quit", command=root.quit,
                   style="Danger.TButton").pack(side="right")
        self.count_label = ttk.Label(bottom, text="", font=("Segoe UI", 9, "italic"))
        self.count_label.pack(side="right", padx=(0, 15))

    # ---------- Styles ----------
    def setup_styles(self):
        style = ttk.Style()
        style.theme_use("clam")
        style.configure("TFrame", background="#f0f4f8")
        style.configure("TLabel", background="#f0f4f8", foreground="#1a2b3c", font=("Segoe UI", 10))
        style.configure("TButton", font=("Segoe UI", 10, "bold"), padding=6)
        style.map("TButton", background=[("active", "#2a7de1")])
        style.configure("Success.TButton", background="#28a745", foreground="white")
        style.map("Success.TButton", background=[("active", "#218838")])
        style.configure("Primary.TButton", background="#2a7de1", foreground="white")
        style.map("Primary.TButton", background=[("active", "#1a5fb4")])
        style.configure("Danger.TButton", background="#dc3545", foreground="white")
        style.map("Danger.TButton", background=[("active", "#c82333")])
        style.configure("Status.TLabel", background="#e9ecef", foreground="#1a2b3c", font=("Segoe UI", 9), padding=4)
        style.configure("Treeview", background="white", fieldbackground="white", font=("Segoe UI", 9))
        style.map("Treeview", background=[("selected", "#2a7de1")])
        style.configure("Treeview.Heading", background="#d9e2ec", font=("Segoe UI", 10, "bold"))

    # ---------- File browsing ----------
    def browse_file(self, var):
        f = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")])
        if f:
            var.set(f)

    # ---------- Processing ----------
    def process_files(self):
        if not self.sponsor_path.get() or not self.att_path.get():
            messagebox.showerror("Error", "Please select both files.")
            return

        try:
            self.status.config(text="Parsing sponsorship file...")
            # id_to_source, email_to_source, phone_to_source
            reg_data = self.parse_sponsorship(self.sponsor_path.get())
            if not reg_data:
                messagebox.showerror("Error", "No sponsorship data found.")
                return

            id_to_source, email_to_source, phone_to_source, email_to_id, phone_to_id = reg_data

            self.status.config(text="Parsing attendance file...")
            attendance = self.parse_attendance(self.att_path.get())
            if not attendance:
                messagebox.showerror("Error", "No attendance data found.")
                return

            matched_count = 0
            unmatched_names = []
            self.result_data = []

            for att in attendance:
                name = att["name"]
                email = att.get("email", "").strip()
                phone = att.get("phone", "").strip()
                idnumber = att.get("idnumber", "").strip()
                source = ""

                # ----- Step 1: If attendance already has an IDNumber -----
                if idnumber:
                    if idnumber in id_to_source:
                        source = id_to_source[idnumber]
                        matched_count += 1
                    else:
                        # ID exists in attendance but not in registration – try fallback
                        # (we still keep the original ID, but try to get source via email/phone)
                        if email and email.lower() in email_to_source:
                            source = email_to_source[email.lower()]
                            matched_count += 1
                        elif phone and clean_phone(phone) in phone_to_source:
                            source = phone_to_source[clean_phone(phone)]
                            matched_count += 1
                        else:
                            unmatched_names.append(name)
                else:
                    # ----- Step 2: No ID, try to get ID and Source via Email or Phone -----
                    if email and email.lower() in email_to_id:
                        idnumber = email_to_id[email.lower()]
                        source = email_to_source[email.lower()]
                        matched_count += 1
                    elif phone and clean_phone(phone) in phone_to_id:
                        idnumber = phone_to_id[clean_phone(phone)]
                        source = phone_to_source[clean_phone(phone)]
                        matched_count += 1
                    else:
                        unmatched_names.append(name)

                row = {
                    "Name": name,
                    "Email": email,
                    "Phone": phone,
                    "Source": source,
                    "IDNumber": idnumber,
                    "Join": att["join"],
                    "Left": att["left"],
                    "Duration": att["duration"]
                }
                self.result_data.append(row)

            self.populate_tree(self.result_data)

            total = len(attendance)
            self.status.config(text=f"✅ Matching complete: {matched_count}/{total} matched. "
                                     f"Unmatched: {len(unmatched_names)}")
            self.count_label.config(text=f"{total} attendance records, {matched_count} have source")

            if unmatched_names:
                msg = "The following attendance records could NOT be matched:\n" + "\n".join(unmatched_names[:20])
                if len(unmatched_names) > 20:
                    msg += f"\n... and {len(unmatched_names)-20} more."
                messagebox.showwarning("Unmatched Names", msg)

        except Exception as e:
            messagebox.showerror("Error", f"Processing failed:\n{str(e)}")
            self.status.config(text="❌ Error")

    # ---------- Parsers ----------
    def parse_sponsorship(self, filepath):
        """
        Returns a tuple:
          (id_to_source, email_to_source, phone_to_source, email_to_id, phone_to_id)
        """
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active

        id_to_source = {}
        email_to_source = {}
        phone_to_source = {}
        email_to_id = {}
        phone_to_id = {}

        # Find header row to map columns
        header_row = None
        col_map = {}
        for row in ws.iter_rows(values_only=True):
            if all(c is None for c in row):
                continue
            row_str = [str(c).strip() if c is not None else "" for c in row]
            found = {}
            for idx, val in enumerate(row_str):
                val_lower = val.lower()
                if "idnumber" in val_lower or "id" in val_lower:
                    found["idnumber"] = idx
                elif "source" in val_lower or "sponsor" in val_lower:
                    found["source"] = idx
                elif "email" in val_lower:
                    found["email"] = idx
                elif "phone" in val_lower:
                    found["phone"] = idx
                elif "full name" in val_lower or "name" in val_lower:
                    found["name"] = idx
            if "idnumber" in found and "source" in found:
                col_map = found
                header_row = row
                break

        if not col_map:
            messagebox.showerror("Error", "Could not find required columns (IDNumber and Source) in sponsorship file.")
            return None

        # Read data rows
        for row in ws.iter_rows(values_only=True):
            if row == header_row:
                continue
            if all(c is None for c in row):
                continue
            if len(row) <= max(col_map.values()):
                continue

            idnumber = str(row[col_map["idnumber"]]).strip() if row[col_map["idnumber"]] else ""
            source = str(row[col_map["source"]]).strip() if row[col_map["source"]] else ""
            email = str(row[col_map.get("email", 0)]).strip() if "email" in col_map and row[col_map["email"]] else ""
            phone_raw = str(row[col_map.get("phone", 0)]).strip() if "phone" in col_map and row[col_map["phone"]] else ""

            if idnumber and source:
                id_to_source[idnumber] = source
                if email:
                    email_key = email.lower()
                    email_to_source[email_key] = source
                    email_to_id[email_key] = idnumber
                if phone_raw:
                    phone_list = parse_phone_list(phone_raw)
                    for p in phone_list:
                        if p:
                            phone_to_source[p] = source
                            phone_to_id[p] = idnumber

        return (id_to_source, email_to_source, phone_to_source, email_to_id, phone_to_id)

    def parse_attendance(self, filepath):
        """Return list of dicts: name, email, phone, idnumber, join, left, duration."""
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active

        header_row = None
        col_map = {}

        # Find header row by scanning for keywords
        for row in ws.iter_rows(values_only=True):
            if all(c is None for c in row):
                continue
            row_str = [str(c).strip() if c is not None else "" for c in row]
            found = {}
            for idx, val in enumerate(row_str):
                val_lower = val.lower()
                if "name" in val_lower:
                    found["name"] = idx
                elif "email" in val_lower:
                    found["email"] = idx
                elif "phone" in val_lower:
                    found["phone"] = idx
                elif "idnumber" in val_lower or "id" in val_lower:
                    found["idnumber"] = idx
                elif "join" in val_lower:
                    found["join"] = idx
                elif "left" in val_lower:
                    found["left"] = idx
                elif "duration" in val_lower:
                    found["duration"] = idx
            if "name" in found and "join" in found and "left" in found and "duration" in found:
                col_map = found
                header_row = row
                break

        if not col_map:
            # Fallback: assume order: Name, Email, Phone, IDNumber, Join, Left, Duration
            first_row = None
            for row in ws.iter_rows(values_only=True):
                if all(c is None for c in row):
                    continue
                first_row = row
                break
            if first_row:
                row_str = [str(c).strip().lower() if c else "" for c in first_row]
                if any("name" in s for s in row_str):
                    col_map = {"name": 0, "email": 1, "phone": 2, "idnumber": 3,
                               "join": 4, "left": 5, "duration": 6}
                    header_row = first_row
                else:
                    col_map = {"name": 0, "email": 1, "phone": 2, "idnumber": 3,
                               "join": 4, "left": 5, "duration": 6}
                    header_row = None

        att_list = []
        for row in ws.iter_rows(values_only=True):
            if header_row is not None and row == header_row:
                continue
            if all(c is None for c in row):
                continue
            if len(row) <= max(col_map.values()):
                continue

            name = str(row[col_map["name"]]).strip() if row[col_map["name"]] else ""
            if not name:
                continue

            email = str(row[col_map.get("email", 0)]).strip() if "email" in col_map and row[col_map["email"]] else ""
            phone = str(row[col_map.get("phone", 0)]).strip() if "phone" in col_map and row[col_map["phone"]] else ""
            idnumber = str(row[col_map.get("idnumber", 0)]).strip() if "idnumber" in col_map and row[col_map["idnumber"]] else ""
            join = row[col_map["join"]]
            left = row[col_map["left"]]
            dur = row[col_map["duration"]]

            if isinstance(dur, (int, float)):
                h = int(dur // 3600)
                m = int((dur % 3600) // 60)
                s = int(dur % 60)
                dur_str = f"{h:02d}:{m:02d}:{s:02d}"
            else:
                dur_str = str(dur).strip()

            att_list.append({
                "name": name,
                "email": email,
                "phone": phone,
                "idnumber": idnumber,
                "join": join,
                "left": left,
                "duration": dur_str
            })

        return att_list

    # ---------- Treeview ----------
    def populate_tree(self, rows):
        for item in self.tree.get_children():
            self.tree.delete(item)
        for row in rows:
            values = (
                row["Name"],
                row["Email"],
                row["Phone"],
                row["Source"],
                row["IDNumber"],
                row["Join"],
                row["Left"],
                row["Duration"]
            )
            self.tree.insert("", "end", values=values)

    # ---------- Save ----------
    def save_result(self):
        if not self.result_data:
            messagebox.showwarning("Warning", "No data to save. Process files first.")
            return

        file_path = filedialog.asksaveasfilename(
            title="Save Enriched Attendance",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        if file_path:
            try:
                wb = Workbook()
                ws = wb.active
                ws.title = "Enriched Attendance"

                headers = ["Name", "Email", "Phone", "Source", "IDNumber", "Join", "Left", "Duration"]
                for col_idx, hdr in enumerate(headers, 1):
                    ws.cell(row=1, column=col_idx, value=hdr)

                for row_idx, row in enumerate(self.result_data, start=2):
                    ws.cell(row=row_idx, column=1, value=row["Name"])
                    ws.cell(row=row_idx, column=2, value=row["Email"])
                    ws.cell(row=row_idx, column=3, value=row["Phone"])
                    ws.cell(row=row_idx, column=4, value=row["Source"])
                    ws.cell(row=row_idx, column=5, value=row["IDNumber"])
                    ws.cell(row=row_idx, column=6, value=row["Join"])
                    ws.cell(row=row_idx, column=7, value=row["Left"])
                    ws.cell(row=row_idx, column=8, value=row["Duration"])

                for col_idx in range(1, 9):
                    col_letter = get_column_letter(col_idx)
                    max_len = max(len(headers[col_idx-1]), 10)
                    for row_idx in range(2, len(self.result_data)+2):
                        val = ws.cell(row=row_idx, column=col_idx).value
                        if val:
                            max_len = max(max_len, len(str(val)))
                    ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

                wb.save(file_path)
                messagebox.showinfo("Success", f"File saved to:\n{file_path}")
                self.status.config(text=f"💾 Saved: {os.path.basename(file_path)}")
            except Exception as e:
                messagebox.showerror("Error", f"Could not save file:\n{str(e)}")

    # ---------- Clear ----------
    def clear_all(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.result_data = []
        self.sponsor_path.set("")
        self.att_path.set("")
        self.status.config(text="Cleared")
        self.count_label.config(text="")

if __name__ == "__main__":
    root = tk.Tk()
    app = SourceAdderApp(root)
    root.mainloop()