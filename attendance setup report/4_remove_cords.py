import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
import csv
import os

# ---------- Custom ttk styles ----------
def setup_styles():
    style = ttk.Style()
    style.theme_use('clam')

    bg = "#f0f4f8"
    fg = "#1a2b3c"
    accent = "#2a7de1"
    success = "#28a745"
    danger = "#dc3545"
    
    style.configure("TFrame", background=bg)
    style.configure("TLabel", background=bg, foreground=fg, font=("Segoe UI", 10))
    style.configure("TButton", font=("Segoe UI", 10, "bold"), padding=6)
    style.map("TButton",
              background=[("active", accent), ("pressed", accent)],
              foreground=[("active", "white"), ("pressed", "white")])

    style.configure("Success.TButton", background=success, foreground="white")
    style.map("Success.TButton",
              background=[("active", "#218838"), ("pressed", "#1e7e34")])

    style.configure("Primary.TButton", background=accent, foreground="white")
    style.map("Primary.TButton",
              background=[("active", "#1a5fb4"), ("pressed", "#15539e")])

    style.configure("Danger.TButton", background=danger, foreground="white")
    style.map("Danger.TButton",
              background=[("active", "#c82333"), ("pressed", "#bd2130")])

    style.configure("Treeview",
                    background="white",
                    foreground=fg,
                    fieldbackground="white",
                    font=("Segoe UI", 9))
    style.map("Treeview", background=[("selected", accent)])

    style.configure("Treeview.Heading",
                    background="#d9e2ec",
                    foreground=fg,
                    font=("Segoe UI", 10, "bold"))

    style.configure("Status.TLabel",
                    background="#e9ecef",
                    foreground=fg,
                    font=("Segoe UI", 9),
                    padding=4)

    return style

# ---------- Main Application ----------
class DataCleanerApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Data Cleaner - Remove Empty/Null Rows")
        self.root.geometry("1050x700")
        self.root.configure(bg="#f0f4f8")

        self.style = setup_styles()

        self.file_path = tk.StringVar()
        self.selected_column = tk.StringVar()
        
        # Storage for processing
        self.headers = []
        self.kept_rows = []
        self.removed_rows = []

        main_frame = ttk.Frame(root, padding="15")
        main_frame.pack(fill="both", expand=True)

        # ----- Header / Title -----
        title_frame = ttk.Frame(main_frame)
        title_frame.pack(fill="x", pady=(0, 10))

        title_label = ttk.Label(title_frame,
                                text="🧹 Data Cleaner & Null Remover",
                                font=("Segoe UI", 18, "bold"),
                                foreground="#1a2b3c")
        title_label.pack(side="left")

        subtitle = ttk.Label(title_frame,
                             text="Preview shows removed rows. Saves a new, clean Excel file.",
                             font=("Segoe UI", 10),
                             foreground="#dc3545")
        subtitle.pack(side="left", padx=(15, 0))

        # ----- File selection & Toolbar -----
        file_frame = ttk.Frame(main_frame)
        file_frame.pack(fill="x", pady=(15, 10))

        ttk.Label(file_frame, text="Data File:").pack(side="left", padx=(0, 8))

        entry = ttk.Entry(file_frame, textvariable=self.file_path, width=35, font=("Segoe UI", 10))
        entry.pack(side="left", padx=(0, 8), fill="x", expand=True)

        browse_btn = ttk.Button(file_frame, text="📂 Browse", command=self.browse_file)
        browse_btn.pack(side="left", padx=(0, 6))

        # Column Selection Dropdown
        ttk.Label(file_frame, text="Check Column:").pack(side="left", padx=(10, 5))
        self.col_combo = ttk.Combobox(file_frame, textvariable=self.selected_column, state="disabled", width=22)
        self.col_combo.pack(side="left", padx=(0, 10))

        process_btn = ttk.Button(file_frame, text="⚡ Process", command=self.process_file, style="Success.TButton")
        process_btn.pack(side="left", padx=(0, 6))

        clear_btn = ttk.Button(file_frame, text="🗑️ Clear", command=self.clear_all, style="Danger.TButton")
        clear_btn.pack(side="left", padx=(0, 6))

        save_btn = ttk.Button(file_frame, text="💾 Save Cleaned Result", command=self.save_result, style="Primary.TButton")
        save_btn.pack(side="left")

        # ----- Status bar -----
        self.status = ttk.Label(main_frame, text="Ready. Please select a file.", relief="sunken",
                                anchor="w", style="Status.TLabel")
        self.status.pack(fill="x", pady=(10, 10))

        # ----- Treeview Header (Removed Rows) -----
        self.tree_title_label = ttk.Label(main_frame, text="⚠️ Removed Rows Preview:", font=("Segoe UI", 10, "bold"), foreground="#dc3545")
        self.tree_title_label.pack(anchor="w", pady=(0, 5))

        # ----- Treeview (Dynamic Columns) -----
        tree_frame = ttk.Frame(main_frame)
        tree_frame.pack(fill="both", expand=True)

        self.tree = ttk.Treeview(tree_frame, show="headings", height=16)
        
        yscrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        xscrollbar = ttk.Scrollbar(tree_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=yscrollbar.set, xscrollcommand=xscrollbar.set)
        
        yscrollbar.pack(side="right", fill="y")
        xscrollbar.pack(side="bottom", fill="x")
        self.tree.pack(side="left", fill="both", expand=True)

        # ----- Bottom action bar -----
        bottom_frame = ttk.Frame(main_frame)
        bottom_frame.pack(fill="x", pady=(12, 0))

        ttk.Button(bottom_frame, text="Quit", command=root.quit, style="Danger.TButton").pack(side="right")

        self.count_label = ttk.Label(bottom_frame, text="", font=("Segoe UI", 9, "italic"))
        self.count_label.pack(side="right", padx=(0, 15))

    # ---------- Methods ----------
    def browse_file(self):
        filename = filedialog.askopenfilename(
            title="Select Data File",
            filetypes=[("Excel & CSV", "*.csv *.xlsx *.xlsm"), ("Excel files", "*.xlsx *.xlsm"), ("CSV files", "*.csv")]
        )
        if filename:
            self.file_path.set(filename)
            self.status.config(text=f"📁 Loaded: {os.path.basename(filename)}")
            self.load_columns(filename)

    def load_columns(self, filename):
        try:
            headers = []
            if filename.lower().endswith(('.xlsx', '.xlsm')):
                wb = load_workbook(filename, read_only=True)
                ws = wb.active
                first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
                headers = [str(h) if h is not None else f"Column {i+1}" for i, h in enumerate(first_row)]
                wb.close()
            elif filename.lower().endswith('.csv'):
                with open(filename, 'r', encoding='utf-8-sig') as f:
                    reader = csv.reader(f)
                    headers = next(reader)
            
            self.col_combo['values'] = ["ALL COLUMNS (Check entire row)"] + headers
            self.col_combo.current(0)
            self.col_combo.config(state="readonly")
            
            self.tree.delete(*self.tree.get_children())
            self.tree["columns"] = []
            self.count_label.config(text="")
            
        except Exception as e:
            messagebox.showerror("Error", f"Could not read file headers:\n{str(e)}")
            self.col_combo.config(state="disabled")

    def is_empty(self, val):
        if val is None:
            return True
        if isinstance(val, str) and val.strip() == "":
            return True
        return False

    def process_file(self):
        if not self.file_path.get() or not self.selected_column.get():
            messagebox.showerror("Error", "Please select a file and a column first.")
            return

        filename = self.file_path.get()
        col_name = self.selected_column.get()

        try:
            self.kept_rows = []
            self.removed_rows = []
            
            if filename.lower().endswith(('.xlsx', '.xlsm')):
                wb = load_workbook(filename, data_only=True)
                ws = wb.active
                
                # Read all data into lists
                data = list(ws.iter_rows(values_only=True))
                if not data:
                    return
                    
                self.headers = [str(cell) if cell is not None else "" for cell in data[0]]
                rows = data[1:]
                
                if col_name == "ALL COLUMNS (Check entire row)":
                    col_idx = None
                else:
                    if col_name not in self.headers:
                        messagebox.showerror("Error", "Column not found.")
                        return
                    col_idx = self.headers.index(col_name)

                # Separate rows
                for row in rows:
                    if col_idx is None:
                        should_remove = any(self.is_empty(v) for v in row)
                    else:
                        val = row[col_idx] if len(row) > col_idx else None
                        should_remove = self.is_empty(val)

                    if should_remove:
                        self.removed_rows.append(row)
                    else:
                        self.kept_rows.append(row)
                wb.close()

            elif filename.lower().endswith('.csv'):
                with open(filename, 'r', encoding='utf-8-sig') as f:
                    reader = csv.reader(f)
                    data = list(reader)
                
                if not data:
                    return
                    
                self.headers = data[0]
                rows = data[1:]
                
                if col_name == "ALL COLUMNS (Check entire row)":
                    col_idx = None
                else:
                    if col_name not in self.headers:
                        messagebox.showerror("Error", "Column not found.")
                        return
                    col_idx = self.headers.index(col_name)

                for row in rows:
                    if col_idx is None:
                        should_remove = any(self.is_empty(v) for v in row)
                    else:
                        val = row[col_idx] if len(row) > col_idx else None
                        should_remove = self.is_empty(val)

                    if should_remove:
                        self.removed_rows.append(row)
                    else:
                        self.kept_rows.append(row)

            # --- Populate Preview ---
            self.populate_tree(self.removed_rows, col_name)
            
            self.status.config(text=f"✅ Done: Removed {len(self.removed_rows)} rows. Remaining: {len(self.kept_rows)}")
            self.count_label.config(text=f"{len(self.removed_rows)} rows removed")
            
            messagebox.showinfo("Summary", 
                                f"Process Complete!\n\n"
                                f"Rows Removed: {len(self.removed_rows)}\n"
                                f"Remaining Rows: {len(self.kept_rows)}\n\n"
                                f"Click 'Save Cleaned Result' to save the remaining rows as a new Excel file.")

        except Exception as e:
            messagebox.showerror("Error", f"Failed to process file:\n{str(e)}")

    def populate_tree(self, removed_rows, col_name):
        self.tree.delete(*self.tree.get_children())
        
        if not removed_rows:
            self.tree["columns"] = ["Status"]
            self.tree.heading("Status", text="No empty rows found!")
            self.tree.column("Status", width=300, anchor="center")
            return

        cols = self.headers
        self.tree["columns"] = cols
        
        for i, col in enumerate(cols):
            self.tree.heading(col, text=col)
            width = 250 if col_name == col else 150
            self.tree.column(col, width=width, anchor="w", stretch=True)
            
        for row in removed_rows[:1000]: # Limit preview to 1000 rows
            clean_row = [str(val) if val is not None else "" for val in row]
            while len(clean_row) < len(cols):
                clean_row.append("")
            self.tree.insert("", "end", values=clean_row)

    # ---------- Save (Using the sample's method) ----------
    def save_result(self):
        if not self.kept_rows:
            messagebox.showwarning("Warning", "No cleaned data to save. Process a file first.")
            return

        file_path = filedialog.asksaveasfilename(
            title="Save Cleaned Data",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        
        if file_path:
            try:
                # Create a brand new workbook
                wb = Workbook()
                ws = wb.active
                ws.title = "Enriched Attendance"

                # Write Headers
                for col_idx, hdr in enumerate(self.headers, 1):
                    ws.cell(row=1, column=col_idx, value=hdr)

                # Write Kept Rows
                for row_idx, row in enumerate(self.kept_rows, start=2):
                    for col_idx, val in enumerate(row, 1):
                        ws.cell(row=row_idx, column=col_idx, value=val)

                # Adjust Column Widths (from your sample)
                for col_idx in range(1, len(self.headers) + 1):
                    col_letter = get_column_letter(col_idx)
                    max_len = max(len(str(self.headers[col_idx-1])), 10)
                    for row_idx in range(2, len(self.kept_rows)+2):
                        val = ws.cell(row=row_idx, column=col_idx).value
                        if val:
                            max_len = max(max_len, len(str(val)))
                    ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

                wb.save(file_path)
                messagebox.showinfo("Success", f"File saved successfully to:\n{file_path}")
                self.status.config(text=f"💾 Saved: {os.path.basename(file_path)}")
            except Exception as e:
                messagebox.showerror("Error", f"Could not save file:\n{str(e)}")

    def clear_all(self):
        self.tree.delete(*self.tree.get_children())
        self.tree["columns"] = []
        self.headers = []
        self.kept_rows = []
        self.removed_rows = []
        self.file_path.set("")
        self.selected_column.set("")
        self.col_combo['values'] = []
        self.col_combo.config(state="disabled")
        self.status.config(text="Cleared")
        self.count_label.config(text="")

if __name__ == "__main__":
    root = tk.Tk()
    app = DataCleanerApp(root)
    root.mainloop()