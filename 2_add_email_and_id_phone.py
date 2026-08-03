import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
import difflib
import re
import os

# ---------- Helper functions ----------
def clean_name(name):
    """Normalize name: lower case, remove punctuation, sort tokens."""
    if not name:
        return ""
    s = str(name).strip()
    s = re.sub(r'[^\w\s]', '', s)
    tokens = s.lower().split()
    tokens = [t for t in tokens if t]
    return " ".join(sorted(tokens))

def parse_duration(value):
    if value is None:
        return 0
    s = str(value).strip()
    parts = s.split(':')
    if len(parts) == 3:
        h, m, sec = map(int, parts)
        return h * 3600 + m * 60 + sec
    elif len(parts) == 2:
        m, sec = map(int, parts)
        return m * 60 + sec
    else:
        try:
            return int(s)
        except:
            return 0

def format_duration(seconds):
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    return f"{h:02d}:{m:02d}:{s:02d}"

# ---------- Main Application ----------
class MergeIDApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Attendance + Registration (ID) Merger")
        self.root.geometry("1200x650")
        self.root.configure(bg="#f0f4f8")

        self.setup_styles()

        self.reg_path = tk.StringVar()
        self.att_path = tk.StringVar()
        self.result_data = []
        self.reg_map = {}
        self.unmatched_names = []

        # ----- Main container -----
        main = ttk.Frame(root, padding="15")
        main.pack(fill="both", expand=True)

        # Title
        ttk.Label(main, text="📋 Attendance + Registration (ID) Merger",
                  font=("Segoe UI", 16, "bold"), foreground="#1a2b3c").pack(anchor="w")
        ttk.Label(main, text="Match attendance names with registration details (email, ID & phone)",
                  font=("Segoe UI", 10), foreground="#5a6b7c").pack(anchor="w", pady=(0, 5))

        # Description
        desc_text = (
            "📌 This app merges your attendance list with registration data.\n"
            "   • Select the Registration file (must contain columns: Name, Email, IDNumber, Phone).\n"
            "   • Select the Attendance file (columns: Name, Join, Left, Duration).\n"
            "   • Click 'Merge & Match' – the app will match by name (handles variations).\n"
            "   • The result will include Email, IDNumber, and Phone from registration.\n"
            "   • Unmatched names are reported – you can then manually correct them."
        )
        desc_label = ttk.Label(main, text=desc_text, font=("Segoe UI", 9), foreground="#2a4b6c",
                               background="#f0f4f8", justify="left", wraplength=900)
        desc_label.pack(anchor="w", pady=(0, 10))

        ttk.Separator(main, orient="horizontal").pack(fill="x", pady=(0, 10))

        # ----- File selection -----
        reg_frame = ttk.Frame(main)
        reg_frame.pack(fill="x", pady=3)
        ttk.Label(reg_frame, text="📁 Registration File:").pack(side="left", padx=(0, 8))
        ttk.Entry(reg_frame, textvariable=self.reg_path, width=60).pack(side="left", fill="x", expand=True, padx=(0, 8))
        ttk.Button(reg_frame, text="Browse", command=lambda: self.browse_file(self.reg_path)).pack(side="left")

        att_frame = ttk.Frame(main)
        att_frame.pack(fill="x", pady=3)
        ttk.Label(att_frame, text="📁 Attendance File:").pack(side="left", padx=(0, 8))
        ttk.Entry(att_frame, textvariable=self.att_path, width=60).pack(side="left", fill="x", expand=True, padx=(0, 8))
        ttk.Button(att_frame, text="Browse", command=lambda: self.browse_file(self.att_path)).pack(side="left")

        # Toolbar with Save button next to Clear
        btn_frame = ttk.Frame(main)
        btn_frame.pack(fill="x", pady=10)

        ttk.Button(btn_frame, text="⚡ Merge & Match", command=self.process_files,
                   style="Success.TButton").pack(side="left", padx=(0, 10))
        ttk.Button(btn_frame, text="Clear All", command=self.clear_all,
                   style="Danger.TButton").pack(side="left", padx=(0, 10))
        ttk.Button(btn_frame, text="💾 Save as Excel", command=self.save_result,
                   style="Primary.TButton").pack(side="left")

        # Status bar
        self.status = ttk.Label(main, text="Ready", relief="sunken", anchor="w", style="Status.TLabel")
        self.status.pack(fill="x", pady=(5, 10))

        # Treeview
        tree_frame = ttk.Frame(main)
        tree_frame.pack(fill="both", expand=True)

        columns = ("Name", "Email", "Phone", "IDNumber", "Join", "Left", "Duration")
        self.tree = ttk.Treeview(tree_frame, columns=columns, show="headings", height=18)

        col_widths = {"Name": 200, "Email": 200, "Phone": 120, "IDNumber": 150, "Join": 150, "Left": 150, "Duration": 100}
        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=col_widths.get(col, 120),
                             anchor="w" if col=="Name" else "center")

        scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        self.tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # Bottom bar (Quit + count label)
        bottom = ttk.Frame(main)
        bottom.pack(fill="x", pady=10)
        ttk.Button(bottom, text="Quit", command=root.quit,
                   style="Danger.TButton").pack(side="right")
        self.count_label = ttk.Label(bottom, text="", font=("Segoe UI", 9, "italic"))
        self.count_label.pack(side="right", padx=(0, 15))

    # ---------- Styles (unchanged) ----------
    def setup_styles(self):
        style = ttk.Style()
        style.theme_use("clam")
        style.configure("TFrame", background="#f0f4f8")
        style.configure("TLabel", background="#f0f4f8", foreground="#1a2b3c", font=("Segoe UI", 10))
        style.configure("TButton", font=("Segoe UI", 10, "bold"), padding=6)
        style.map("TButton", background=[("active", "#2a7de1")])
        style.configure("Success.TButton", background="#28a745", foreground="white")
        style.map("Success.TButton", background=[("active", "#218838")])
        style.configure("Primary.TButton", background="#2a7de1", foreground="white")
        style.map("Primary.TButton", background=[("active", "#1a5fb4")])
        style.configure("Danger.TButton", background="#dc3545", foreground="white")
        style.map("Danger.TButton", background=[("active", "#c82333")])
        style.configure("Status.TLabel", background="#e9ecef", foreground="#1a2b3c", font=("Segoe UI", 9), padding=4)
        style.configure("Treeview", background="white", fieldbackground="white", font=("Segoe UI", 9))
        style.map("Treeview", background=[("selected", "#2a7de1")])
        style.configure("Treeview.Heading", background="#d9e2ec", font=("Segoe UI", 10, "bold"))

    # ---------- File browsing ----------
    def browse_file(self, var):
        f = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")])
        if f:
            var.set(f)

    # ---------- Processing (unchanged) ----------
    def process_files(self):
        if not self.reg_path.get() or not self.att_path.get():
            messagebox.showerror("Error", "Please select both files.")
            return

        try:
            self.status.config(text="Parsing registration file...")
            reg_data = self.parse_registration(self.reg_path.get())
            if not reg_data:
                messagebox.showerror("Error", "No registration data found.")
                return

            self.reg_map = {}
            for item in reg_data:
                cleaned = clean_name(item["name"])
                if cleaned not in self.reg_map:
                    self.reg_map[cleaned] = (item["email"], item["idnumber"], item.get("phone", ""))

            self.status.config(text="Parsing attendance file...")
            attendance = self.parse_attendance(self.att_path.get())
            if not attendance:
                messagebox.showerror("Error", "No attendance data found.")
                return

            matched_count = 0
            self.unmatched_names = []
            self.result_data = []

            for att in attendance:
                name = att["name"]
                cleaned = clean_name(name)
                email, idnumber, phone = "", "", ""

                if cleaned in self.reg_map:
                    email, idnumber, phone = self.reg_map[cleaned]
                    matched_count += 1
                else:
                    candidates = difflib.get_close_matches(cleaned, self.reg_map.keys(), n=1, cutoff=0.7)
                    if candidates:
                        email, idnumber, phone = self.reg_map[candidates[0]]
                        matched_count += 1
                    else:
                        self.unmatched_names.append(name)

                row = {
                    "Name": name,
                    "Email": email,
                    "Phone": phone,
                    "IDNumber": idnumber,
                    "Join": att["join"],
                    "Left": att["left"],
                    "Duration": format_duration(att["duration_seconds"])
                }
                self.result_data.append(row)

            self.populate_tree(self.result_data)

            total = len(attendance)
            self.status.config(text=f"✅ Matching complete: {matched_count}/{total} matched. "
                                     f"Unmatched: {len(self.unmatched_names)}")
            self.count_label.config(text=f"{total} attendance records, {matched_count} have registration")

            if self.unmatched_names:
                msg = "The following attendance names could NOT be matched:\n" + "\n".join(self.unmatched_names[:20])
                if len(self.unmatched_names) > 20:
                    msg += f"\n... and {len(self.unmatched_names)-20} more."
                messagebox.showwarning("Unmatched Names", msg)

        except Exception as e:
            messagebox.showerror("Error", f"Processing failed:\n{str(e)}")
            self.status.config(text="❌ Error")

    # ---------- Parsers (unchanged) ----------
    def parse_registration(self, filepath):
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active
        data = []
        header_row = None
        col_map = {}

        for row in ws.iter_rows(values_only=True):
            if all(c is None for c in row):
                continue
            row_str = [str(c).strip() if c is not None else "" for c in row]
            found = {}
            for idx, val in enumerate(row_str):
                val_lower = val.lower()
                if "email" in val_lower:
                    found["email"] = idx
                elif "name" in val_lower:
                    found["name"] = idx
                elif "idnumber" in val_lower or "id" in val_lower:
                    found["idnumber"] = idx
                elif "phone" in val_lower:
                    found["phone"] = idx
            if "email" in found and "name" in found and "idnumber" in found:
                col_map = found
                header_row = row
                break

        if not col_map:
            first_row = None
            for row in ws.iter_rows(values_only=True):
                if all(c is None for c in row):
                    continue
                first_row = row
                break
            if first_row and len(first_row) >= 4:
                email_col = None
                name_col = None
                id_col = None
                phone_col = None
                for idx, val in enumerate(first_row):
                    if val:
                        val_lower = str(val).lower()
                        if '@' in val_lower:
                            email_col = idx
                        elif 'id' in val_lower:
                            id_col = idx
                        elif 'phone' in val_lower:
                            phone_col = idx
                if email_col is None:
                    email_col = 1
                if id_col is None:
                    id_col = 3
                if phone_col is None:
                    phone_col = 4 if len(first_row) > 4 else None
                name_col = None
                for idx in range(len(first_row)):
                    if idx not in (email_col, id_col, phone_col) if phone_col is not None else (email_col, id_col):
                        name_col = idx
                        break
                if name_col is None:
                    name_col = 2
                col_map = {"email": email_col, "name": name_col, "idnumber": id_col}
                if phone_col is not None:
                    col_map["phone"] = phone_col
                header_row = first_row
            else:
                return []

        for row in ws.iter_rows(values_only=True):
            if row == header_row:
                continue
            if all(c is None for c in row):
                continue
            if len(row) <= max(col_map.values()):
                continue
            name = str(row[col_map["name"]]).strip() if row[col_map["name"]] else ""
            email = str(row[col_map["email"]]).strip() if row[col_map["email"]] else ""
            idnumber = str(row[col_map["idnumber"]]).strip() if row[col_map["idnumber"]] else ""
            phone = ""
            if "phone" in col_map:
                phone = str(row[col_map["phone"]]).strip() if row[col_map["phone"]] else ""
            if name:
                data.append({"name": name, "email": email, "idnumber": idnumber, "phone": phone})
        return data

    def parse_attendance(self, filepath):
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active
        att_list = []
        first_row = None
        for row in ws.iter_rows(values_only=True):
            if all(c is None for c in row):
                continue
            first_row = row
            break
        if first_row is None:
            return []

        header = False
        if first_row and len(first_row) >= 4:
            if any("name" in str(c).lower() for c in first_row if c):
                header = True
            elif any("join" in str(c).lower() for c in first_row if c):
                header = True

        start_row = 1 if header else 0
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if row_idx < start_row:
                continue
            if all(c is None for c in row):
                continue
            if len(row) < 4:
                continue
            name = str(row[0]).strip() if row[0] else ""
            if not name:
                continue
            join = row[1]
            left = row[2]
            dur = row[3]
            dur_sec = parse_duration(dur)
            att_list.append({
                "name": name,
                "join": join,
                "left": left,
                "duration_seconds": dur_sec
            })
        return att_list

    # ---------- Treeview ----------
    def populate_tree(self, rows):
        for item in self.tree.get_children():
            self.tree.delete(item)
        for row in rows:
            values = (
                row["Name"],
                row["Email"],
                row["Phone"],
                row["IDNumber"],
                row["Join"],
                row["Left"],
                row["Duration"]
            )
            self.tree.insert("", "end", values=values)

    # ---------- Save ----------
    def save_result(self):
        if not self.result_data:
            messagebox.showwarning("Warning", "No data to save. Process files first.")
            return

        file_path = filedialog.asksaveasfilename(
            title="Save Merged Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        if file_path:
            try:
                wb = Workbook()
                ws = wb.active
                ws.title = "Merged"

                headers = ["Name", "Email", "Phone", "IDNumber", "Join", "Left", "Duration"]
                for col_idx, hdr in enumerate(headers, 1):
                    ws.cell(row=1, column=col_idx, value=hdr)

                for row_idx, row in enumerate(self.result_data, start=2):
                    ws.cell(row=row_idx, column=1, value=row["Name"])
                    ws.cell(row=row_idx, column=2, value=row["Email"])
                    ws.cell(row=row_idx, column=3, value=row["Phone"])
                    ws.cell(row=row_idx, column=4, value=row["IDNumber"])
                    ws.cell(row=row_idx, column=5, value=row["Join"])
                    ws.cell(row=row_idx, column=6, value=row["Left"])
                    ws.cell(row=row_idx, column=7, value=row["Duration"])

                for col_idx in range(1, 8):
                    col_letter = get_column_letter(col_idx)
                    max_len = max(len(headers[col_idx-1]), 10)
                    for row_idx in range(2, len(self.result_data)+2):
                        val = ws.cell(row=row_idx, column=col_idx).value
                        if val:
                            max_len = max(max_len, len(str(val)))
                    ws.column_dimensions[col_letter].width = min(max_len + 2, 30)

                wb.save(file_path)
                messagebox.showinfo("Success", f"File saved to:\n{file_path}")
                self.status.config(text=f"💾 Saved: {os.path.basename(file_path)}")
            except Exception as e:
                messagebox.showerror("Error", f"Could not save file:\n{str(e)}")

    # ---------- Clear ----------
    def clear_all(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.result_data = []
        self.reg_map = {}
        self.unmatched_names = []
        self.reg_path.set("")
        self.att_path.set("")
        self.status.config(text="Cleared")
        self.count_label.config(text="")

if __name__ == "__main__":
    root = tk.Tk()
    app = MergeIDApp(root)
    root.mainloop()